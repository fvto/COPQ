"""
Quality Sentinel & Anomaly Triage Agent (Workflow 2)
====================================================
Solves Inefficiency 2:
- Eliminates passive quality checks and silent defect/chart omissions.
- Pre-flight site completeness audit across all 7 manufacturing plants.
- Defect Dictionary & Color Palette coverage check against BC_Color/Color_Defect.xlsx
  to prevent silent PPTX chart omissions.
- Multi-dimensional anomaly audit: Working Hours > 15h, Cost > $1000, negative quantities.
- Generates actionable Quality Audit Reports in Markdown and JSON.
"""

import os
import sys
import re
import json
import logging
from datetime import datetime

import pandas as pd
import numpy as np
import openpyxl

logger = logging.getLogger("COPQ.QualitySentinel")
if not logger.handlers:
    handler = logging.StreamHandler(sys.stdout)
    handler.setFormatter(logging.Formatter("[%(asctime)s] [%(levelname)s] [Sentinel] %(message)s", "%H:%M:%S"))
    logger.addHandler(handler)
    logger.setLevel(logging.INFO)

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from pipeline_common import extract_site_name
from monthly_reporting import classify_copq_type


class QualitySentinelAgent:
    """Proactive Quality Assurance and Data Anomaly Audit Agent."""

    EXPECTED_COPQ_SITES = ["VH", "VH2", "VH3", "VH4", "JV", "JV2", "JV3"]
    KNOWN_FTT_SITES = ["VH", "VH2", "JV", "JV2"]
    
    # Anomaly thresholds
    MAX_WORKING_HOURS = 15.0
    HIGH_COST_THRESHOLD = 1000.0
    MIN_QTY_HIGH_HOURS = 2000

    def __init__(self, root_dir=ROOT_DIR):
        self.root_dir = root_dir
        self.copq_input_dir = os.path.join(root_dir, "COPQ_Input")
        self.ftt_input_dir = os.path.join(root_dir, "FTT_Input")
        self.output_dir = os.path.join(root_dir, "Output")
        self.color_map_file = os.path.join(root_dir, "BC_Color", "Color_Defect.xlsx")
        self.copq_clean_file = os.path.join(self.output_dir, "COPQ_Clean.xlsx")

    def run_full_audit(self, period_label=None):
        """
        Executes a comprehensive pre-flight & post-flight data quality audit.
        """
        logger.info(">>> Starting Proactive Quality & Anomaly Audit...")
        findings = {
            "timestamp": datetime.now().isoformat(),
            "period": period_label or "Current",
            "health_score": 100,
            "status": "PASSED",
            "site_coverage": {},
            "schema_integrity": {},
            "defect_dictionary": {},
            "statistical_anomalies": {},
            "action_items": []
        }

        # 1. Site Coverage Audit
        self._audit_site_coverage(findings)

        # 2. Defect Color Mapping Dictionary Audit
        self._audit_defect_color_mapping(findings)

        # 3. Data Integrity & Anomaly Audit on Combined Data
        self._audit_data_anomalies(findings)

        # Calculate Final Health Score
        deductions = 0
        for item in findings["action_items"]:
            if item["severity"] == "CRITICAL":
                deductions += 25
            elif item["severity"] == "WARNING":
                deductions += 10
            elif item["severity"] == "INFO":
                deductions += 2

        findings["health_score"] = max(0, 100 - deductions)
        if findings["health_score"] < 70:
            findings["status"] = "ACTION_REQUIRED"
        elif findings["health_score"] < 90:
            findings["status"] = "WARNINGS_FOUND"
        else:
            findings["status"] = "HEALTHY"

        logger.info(f"Audit Complete. Health Score: {findings['health_score']}/100 ({findings['status']})")
        
        # Save Reports
        self._export_reports(findings)
        return findings

    def _audit_site_coverage(self, findings):
        logger.info("Auditing Plant Site Coverage...")
        copq_files = [f for f in os.listdir(self.copq_input_dir) if f.endswith((".xls", ".xlsx", ".csv")) and not f.startswith("~$")]
        detected_copq_sites = set()
        for f in copq_files:
            site = extract_site_name(f)
            if site:
                detected_copq_sites.add(site)

        missing_copq = [s for s in self.EXPECTED_COPQ_SITES if s not in detected_copq_sites and not (s == "JV3" and "JVB" in detected_copq_sites)]
        findings["site_coverage"]["copq"] = {
            "total_files": len(copq_files),
            "detected_sites": sorted(list(detected_copq_sites)),
            "missing_sites": missing_copq,
            "is_complete": len(missing_copq) == 0
        }

        if missing_copq:
            findings["action_items"].append({
                "severity": "CRITICAL" if len(missing_copq) > 2 else "WARNING",
                "category": "Site Coverage",
                "message": f"Missing COPQ ERP export files for sites: {', '.join(missing_copq)}. Executive rollups will be incomplete."
            })

        ftt_files = [f for f in os.listdir(self.ftt_input_dir) if f.endswith((".xls", ".xlsx", ".csv")) and not f.startswith("~$")]
        detected_ftt_sites = set()
        for f in ftt_files:
            site = extract_site_name(f)
            if site:
                detected_ftt_sites.add(site)

        findings["site_coverage"]["ftt"] = {
            "total_files": len(ftt_files),
            "detected_sites": sorted(list(detected_ftt_sites)),
            "standard_sites_missing": [s for s in self.KNOWN_FTT_SITES if s not in detected_ftt_sites]
        }

        if len(detected_ftt_sites) < 4:
            findings["action_items"].append({
                "severity": "WARNING",
                "category": "FTT Coverage",
                "message": f"Expected 4 FTT internal reports (VH, VH2, JV, JV2), found {len(detected_ftt_sites)}. Missing: {', '.join([s for s in self.KNOWN_FTT_SITES if s not in detected_ftt_sites])}."
            })

    def _audit_defect_color_mapping(self, findings):
        logger.info("Auditing Defect Dictionary against Color_Defect.xlsx...")
        if not os.path.exists(self.color_map_file):
            findings["action_items"].append({
                "severity": "CRITICAL",
                "category": "Defect Dictionary",
                "message": f"Defect color mapping workbook not found at {self.color_map_file}"
            })
            return

        known_defects = set()
        try:
            wb = openpyxl.load_workbook(self.color_map_file, data_only=True)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for r in range(2, ws.max_row + 1):
                    val = ws.cell(r, 1).value
                    if val:
                        norm = re.sub(r"\s+", "", str(val).strip().lower())
                        known_defects.add(norm)
        except Exception as e:
            logger.warning(f"Could not load color map file: {e}")

        # Scan active FTT raw files for defects
        unmapped_defects = {}
        total_ftt_defects_found = set()
        for fname in os.listdir(self.ftt_input_dir):
            if not fname.endswith((".xlsx", ".xls", ".csv")) or fname.startswith("~$"):
                continue
            fpath = os.path.join(self.ftt_input_dir, fname)
            try:
                df = pd.read_excel(fpath) if not fname.endswith(".csv") else pd.read_csv(fpath, encoding="latin1")
                # Look for defect column
                defect_col = None
                for col in df.columns:
                    c_norm = str(col).lower().strip()
                    if ("defect" in c_norm or "lỗi" in c_norm) and "location" not in c_norm and "position" not in c_norm:
                        defect_col = col
                        break
                if defect_col:
                    raw_defects = df[defect_col].dropna().unique()
                    for d in raw_defects:
                        d_str = str(d).strip()
                        total_ftt_defects_found.add(d_str)
                        d_norm = re.sub(r"\s+", "", d_str.lower())
                        if d_norm in known_defects:
                            continue
                        clean_norm = re.sub(r"^[bc]/", "", d_norm)
                        if clean_norm in known_defects:
                            continue
                        matched_cat = False
                        for k in known_defects:
                            k_clean = re.sub(r"^[bc]/", "", k)
                            if len(k_clean) >= 4 and (k_clean in clean_norm or clean_norm in k_clean):
                                matched_cat = True
                                break
                        if not matched_cat:
                            unmapped_defects[d_str] = unmapped_defects.get(d_str, 0) + 1
            except Exception:
                pass

        findings["defect_dictionary"] = {
            "total_known_palette_defects": len(known_defects),
            "active_ftt_defects_count": len(total_ftt_defects_found),
            "unmapped_defects": sorted(list(unmapped_defects.keys()))
        }

        if unmapped_defects:
            unmapped_list = sorted(list(unmapped_defects.keys()))
            findings["action_items"].append({
                "severity": "WARNING",
                "category": "Defect Dictionary",
                "message": f"Found {len(unmapped_defects)} defect(s) in active FTT logs not present in Color_Defect.xlsx: {', '.join(unmapped_list[:5])}{'...' if len(unmapped_list) > 5 else ''}. These may trigger fallback colors in PPTX charts."
            })

    def _audit_data_anomalies(self, findings):
        logger.info("Auditing Data Anomalies in Master COPQ Records...")
        if not os.path.exists(self.copq_clean_file):
            findings["action_items"].append({
                "severity": "INFO",
                "category": "Data Audit",
                "message": "COPQ_Clean.xlsx not yet generated; run Pipeline Orchestrator to generate master records."
            })
            return

        try:
            df = pd.read_excel(self.copq_clean_file, sheet_name="COPQ_Clean")
        except Exception as e:
            logger.warning(f"Could not read COPQ_Clean sheet: {e}")
            return

        # Anomaly 1: Excessive Working Hours
        high_hours_df = df[df["Working hours"] > self.MAX_WORKING_HOURS] if "Working hours" in df.columns else pd.DataFrame()
        # Anomaly 2: Extreme Cost (> $1000)
        high_cost_df = df[df["Ttl cost ($)"] > self.HIGH_COST_THRESHOLD] if "Ttl cost ($)" in df.columns else pd.DataFrame()
        # Anomaly 3: Negative Values
        neg_cost = df[df["Ttl cost ($)"] < 0] if "Ttl cost ($)" in df.columns else pd.DataFrame()
        neg_qty = df[df["Defective Qty(Pair)"] < 0] if "Defective Qty(Pair)" in df.columns else pd.DataFrame()

        # Anomaly 1B: Working hours > 5 & Defective Qty < 300
        disprop_hours_df = (
            df[(df["Working hours"] > 5.0) & (df["Defective Qty(Pair)"] < 300)]
            if ("Working hours" in df.columns and "Defective Qty(Pair)" in df.columns)
            else pd.DataFrame()
        )
        # Validation error rows from COPQ_Clean
        val_error_df = (
            df[df["Validation Errors"].notna() & (df["Validation Errors"].astype(str).str.strip() != "")]
            if "Validation Errors" in df.columns
            else pd.DataFrame()
        )

        findings["statistical_anomalies"] = {
            "total_records_audited": len(df),
            "high_working_hours_count": len(high_hours_df),
            "disproportionate_hours_count": len(disprop_hours_df),
            "validation_errors_count": len(val_error_df),
            "high_cost_outliers_count": len(high_cost_df),
            "negative_cost_records": len(neg_cost),
            "negative_qty_records": len(neg_qty)
        }

        if len(disprop_hours_df) > 0:
            sample_models = disprop_hours_df["Model"].dropna().head(3).tolist() if "Model" in disprop_hours_df.columns else []
            findings["action_items"].append({
                "severity": "WARNING",
                "category": "Disproportionate Working Hours",
                "message": f"{len(disprop_hours_df)} records have high working hours (> 5h) for low defective qty (< 300 pairs). Samples: {', '.join(map(str, sample_models))}."
            })

        if len(val_error_df) > 0:
            err_samples = val_error_df["Validation Errors"].unique()[:3].tolist()
            findings["action_items"].append({
                "severity": "WARNING",
                "category": "Validation Errors",
                "message": f"{len(val_error_df)} records flagged with validation issues. Samples: {'; '.join(map(str, err_samples))}."
            })

        if len(high_hours_df) > 0:
            sample_models = high_hours_df["Model"].dropna().head(3).tolist() if "Model" in high_hours_df.columns else []
            findings["action_items"].append({
                "severity": "WARNING",
                "category": "Working Hours Anomaly",
                "message": f"{len(high_hours_df)} records exceed standard working hours limit (> {self.MAX_WORKING_HOURS}h). Samples: {', '.join(map(str, sample_models))}."
            })

        if len(high_cost_df) > 0:
            max_cost = high_cost_df["Ttl cost ($)"].max()
            findings["action_items"].append({
                "severity": "INFO",
                "category": "Cost Outliers",
                "message": f"{len(high_cost_df)} records exceed ${self.HIGH_COST_THRESHOLD:,.0f} threshold. Max single incident: ${max_cost:,.2f}."
            })

        if len(neg_cost) > 0 or len(neg_qty) > 0:
            findings["action_items"].append({
                "severity": "CRITICAL",
                "category": "Negative Values",
                "message": f"Found negative values in dataset: {len(neg_cost)} negative cost rows, {len(neg_qty)} negative qty rows."
            })

    def _export_reports(self, findings):
        os.makedirs(self.output_dir, exist_ok=True)
        json_path = os.path.join(self.output_dir, "COPQ_Quality_Audit_Report.json")
        md_path = os.path.join(self.output_dir, "COPQ_Quality_Audit_Report.md")

        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(findings, f, indent=2)

        # Markdown Report Generation
        md_content = [
            f"# COPQ Quality Health & Anomaly Audit Report",
            f"**Audit Timestamp:** {findings['timestamp']}  ",
            f"**Reporting Period:** {findings['period']}  ",
            f"**Overall Health Score:** `{findings['health_score']}/100`  ",
            f"**System Status:** `{findings['status']}`  ",
            "",
            "---",
            "",
            "## 1. Executive Summary & Action Items",
            ""
        ]

        if not findings["action_items"]:
            md_content.append("✅ **All quality checks passed.** No anomalies or missing files detected.")
        else:
            md_content.append("| Severity | Category | Operational Finding / Recommendation |")
            md_content.append("|---|---|---|")
            for item in findings["action_items"]:
                sev_icon = "🔴" if item["severity"] == "CRITICAL" else ("🟡" if item["severity"] == "WARNING" else "ℹ️")
                md_content.append(f"| {sev_icon} **{item['severity']}** | {item['category']} | {item['message']} |")

        md_content.extend([
            "",
            "## 2. Multi-Plant Site Coverage",
            "",
            f"- **COPQ ERP Exports:** {len(findings['site_coverage'].get('copq', {}).get('detected_sites', []))} sites detected (`{', '.join(findings['site_coverage'].get('copq', {}).get('detected_sites', []))}`).",
            f"- **Missing COPQ Sites:** {', '.join(findings['site_coverage'].get('copq', {}).get('missing_sites', [])) or 'None (100% Complete)'}.",
            f"- **FTT B/C Defect Reports:** {len(findings['site_coverage'].get('ftt', {}).get('detected_sites', []))} sites detected (`{', '.join(findings['site_coverage'].get('ftt', {}).get('detected_sites', []))}`).",
            "",
            "## 3. Defect Dictionary & Color Mapping Integrity",
            "",
            f"- **Palette Defect Terms in SOP:** {findings['defect_dictionary'].get('total_known_palette_defects', 0)}",
            f"- **Active Defects Extracted:** {findings['defect_dictionary'].get('active_ftt_defects_count', 0)}",
            f"- **Unmapped Defects:** {len(findings['defect_dictionary'].get('unmapped_defects', []))}"
        ])

        if findings['defect_dictionary'].get('unmapped_defects'):
            md_content.append("")
            md_content.append("**Unmapped Defect Terms:**")
            for d in findings['defect_dictionary']['unmapped_defects']:
                md_content.append(f"- `{d}` -> *Assigned fallback palette color in openpyxl/pptx*")

        md_content.extend([
            "",
            "## 4. Statistical & Heuristic Anomaly Audit",
            "",
            f"- **Total Master Records Audited:** {findings['statistical_anomalies'].get('total_records_audited', 0):,}",
            f"- **High Working Hours (>15h):** {findings['statistical_anomalies'].get('high_working_hours_count', 0)} rows",
            f"- **High Single-Event Cost (>$1,000):** {findings['statistical_anomalies'].get('high_cost_outliers_count', 0)} rows",
            f"- **Negative Quantities / Cost:** {findings['statistical_anomalies'].get('negative_qty_records', 0) + findings['statistical_anomalies'].get('negative_cost_records', 0)} rows",
            "",
            "---",
            "*Report generated proactively by QualitySentinelAgent.*"
        ])

        with open(md_path, "w", encoding="utf-8") as f:
            f.write("\n".join(md_content))

        logger.info(f"Audit reports written to:\n  - {md_path}\n  - {json_path}")
