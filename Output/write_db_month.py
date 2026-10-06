import os
import sys
import copy
import re
import datetime as dt
import openpyxl
from openpyxl.utils import get_column_letter
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
if os.path.basename(HERE) == "Output":
    BASE = os.path.dirname(HERE)
else:
    BASE = HERE

for p in (BASE, os.path.join(BASE, "Source")):
    if p not in sys.path:
        sys.path.insert(0, p)
from monthly_reporting import classify_copq_type

# Default files
CLEAN_XLSX = os.path.join(BASE, "Output", "COPQ_Clean.xlsx")
DB_OVERVIEW = os.path.join(BASE, "Database", "CoPQ database 25.xlsx")
DB_TYPE = os.path.join(BASE, "Database", "CoPQ_type_analysis.xlsx")

# Monthly templates/archives for fallback
CLEAN_BASELINE_TYPE = os.path.join(BASE, "Monthly", "2026-07", "Database", "CoPQ_type_analysis.xlsx")
CLEAN_BASELINE_DB25 = os.path.join(BASE, "Monthly", "2026-07", "Database", "CoPQ database 25.xlsx")


def detect_report_month(clean_path):
    fname = os.path.basename(clean_path)
    month_names = {
        'jan': 1, 'feb': 2, 'mar': 3, 'apr': 4, 'may': 5, 'jun': 6,
        'jul': 7, 'aug': 8, 'sep': 9, 'oct': 10, 'nov': 11, 'dec': 12
    }
    m = re.search(r'clean_([a-z]{3})', fname, re.IGNORECASE)
    if m:
        m_str = m.group(1).lower()
        if m_str in month_names:
            return dt.date(2026, month_names[m_str], 1)
    try:
        df = pd.read_excel(clean_path, sheet_name='COPQ_Clean', nrows=500)
        if 'Date' in df.columns:
            dates = pd.to_datetime(df['Date'], errors='coerce').dropna()
            if not dates.empty:
                max_d = dates.max()
                return dt.date(max_d.year, max_d.month, 1)
    except Exception:
        pass
    return dt.date(2026, 8, 1)


def parse_from_clean_df(clean_path):
    """
    Directly parse COPQ_Clean master sheet using standard SOP classification.
    Produces all necessary structures for both Database workbooks without depending
    on pre-existing pivot sheets.
    """
    df = pd.read_excel(clean_path, sheet_name='COPQ_Clean')
    if "Ttl cost ($)" not in df.columns and "Cost" in df.columns:
        df["Ttl cost ($)"] = pd.to_numeric(df["Cost"], errors="coerce").fillna(0.0)
    if "Defective Qty(Pair)" not in df.columns and "Qty" in df.columns:
        df["Defective Qty(Pair)"] = pd.to_numeric(df["Qty"], errors="coerce").fillna(0.0)
    if "Working hours" not in df.columns:
        df["Working hours"] = 0.0

    df['_group'] = df.apply(classify_copq_type, axis=1)
    df['Site_M'] = df['Site'].astype(str).str.strip().str.upper().replace({'JV3': 'JVB'})

    # 1. top_table
    top_table = {}
    for s in ["VH", "VH2", "VH3", "VH4", "JV", "JV2", "JVB"]:
        sub = df[df['Site_M'] == s]
        tu = round(float(sub[sub['_group'] == 'Touch-up']['Ttl cost ($)'].sum()), 2)
        re_ = round(float(sub[sub['_group'] == 'Reinspection']['Ttl cost ($)'].sum()), 2)
        bc = round(float(sub[sub['_group'] == 'B/C']['Ttl cost ($)'].sum()), 2)
        rw = round(float(sub[sub['_group'] == 'Rework']['Ttl cost ($)'].sum()), 2)
        total = round(tu + re_ + bc + rw, 2)
        re_tu = round(re_ + tu, 2)
        top_table[s] = {
            'tu': tu if tu > 0 else None,
            're': re_ if re_ > 0 else None,
            'bc': bc if bc > 0 else None,
            'rw': rw if rw > 0 else None,
            'total': total if total > 0 else None,
            're_tu': re_tu if re_tu > 0 else None,
        }

    # 2. blocks
    blocks = {
        'Rework': {'left': {}, 'right': {}},
        'Re-Inspection': {'left': {}, 'right': {}},
        'Touch-up': {'left': {}, 'right': {}},
    }
    for s in ["JV2", "JV", "VH2", "VH"]:
        sub = df[df['Site_M'] == s]
        sub_rw = sub[sub['_group'] == 'Rework']
        rw_lines = len(sub_rw['Fty/Line'].dropna().unique()) if not sub_rw.empty else None
        blocks['Rework']['left'][s] = rw_lines if rw_lines else None

        sub_bc = sub[sub['_group'] == 'B/C']
        bc_q = float(sub_bc['Defective Qty(Pair)'].sum())
        blocks['Rework']['right'][s] = round(bc_q, 1) if bc_q > 0 else None

    for s in ["JVB", "JV2", "JV", "VH3", "VH2", "VH"]:
        sub = df[df['Site_M'] == s]
        sub_re = sub[sub['_group'] == 'Reinspection']
        re_q = float(sub_re['Defective Qty(Pair)'].sum())
        re_h = float(sub_re['Working hours'].sum())
        blocks['Re-Inspection']['left'][s] = int(re_q) if re_q > 0 else None
        blocks['Re-Inspection']['right'][s] = round(re_h, 1) if re_h > 0 else None

    for s in ["JVB", "JV", "VH4", "VH2", "VH"]:
        sub = df[df['Site_M'] == s]
        sub_tu = sub[sub['_group'] == 'Touch-up']
        tu_q = float(sub_tu['Defective Qty(Pair)'].sum())
        tu_h = float(sub_tu['Working hours'].sum())
        blocks['Touch-up']['left'][s] = int(tu_q) if tu_q > 0 else None
        blocks['Touch-up']['right'][s] = round(tu_h, 1) if tu_h > 0 else None

    # 3. ov_data
    ov_data = {}
    for s in ["VH", "VH2", "VH3", "VH4", "JV", "JV2", "JVB"]:
        sub = df[df['Site_M'] == s]
        sub_bc = sub[sub['_group'] == 'B/C']
        sub_re = sub[sub['_group'] == 'Reinspection']
        sub_tu = sub[sub['_group'] == 'Touch-up']
        ov_data[s] = {
            'bc_q': float(sub_bc['Defective Qty(Pair)'].sum()) if not sub_bc.empty else None,
            'bc_c': float(sub_bc['Ttl cost ($)'].sum()) if not sub_bc.empty else None,
            're_q': float(sub_re['Defective Qty(Pair)'].sum()) if not sub_re.empty else None,
            're_c': float(sub_re['Ttl cost ($)'].sum()) if not sub_re.empty else None,
            'tu_q': float(sub_tu['Defective Qty(Pair)'].sum()) if not sub_tu.empty else None,
            'tu_c': float(sub_tu['Ttl cost ($)'].sum()) if not sub_tu.empty else None,
        }

    # 4. fac_data
    fac_data = {'BC': {}, 'RE': {}, 'TU': {}}
    for s in ["VH", "VH2", "VH3", "VH4", "JV", "JV2", "JVB"]:
        sub = df[df['Site_M'] == s]
        # BC
        sub_bc = sub[sub['_group'] == 'B/C']
        if not sub_bc.empty:
            g = sub_bc.groupby('Factory').agg(qty=('Defective Qty(Pair)', 'sum'), cost=('Ttl cost ($)', 'sum')).reset_index()
            fac_data['BC'][s] = [{'factory': str(r['Factory']).strip(), 'qty': float(r['qty']), 'cost': float(r['cost'])} for _, r in g.iterrows()]
        # RE
        sub_re = sub[sub['_group'] == 'Reinspection']
        if not sub_re.empty:
            g = sub_re.groupby('Factory').agg(qty=('Defective Qty(Pair)', 'sum'), hrs=('Working hours', 'sum'), cost=('Ttl cost ($)', 'sum')).reset_index()
            fac_data['RE'][s] = [{'factory': str(r['Factory']).strip(), 'qty': float(r['qty']), 'hrs': float(r['hrs']), 'cost': float(r['cost'])} for _, r in g.iterrows()]
        # TU
        sub_tu = sub[sub['_group'] == 'Touch-up']
        if not sub_tu.empty:
            g = sub_tu.groupby('Factory').agg(qty=('Defective Qty(Pair)', 'sum'), hrs=('Working hours', 'sum'), cost=('Ttl cost ($)', 'sum')).reset_index()
            fac_data['TU'][s] = [{'factory': str(r['Factory']).strip(), 'qty': float(r['qty']), 'hrs': float(r['hrs']), 'cost': float(r['cost'])} for _, r in g.iterrows()]

    # 5. mod_data
    mod_data = {'BC': {}, 'RE': {}, 'TU': {}}
    for s in ["VH", "VH2", "VH3", "VH4", "JV", "JV2", "JVB"]:
        sub = df[df['Site_M'] == s]
        # BC
        sub_bc = sub[sub['_group'] == 'B/C']
        if not sub_bc.empty:
            g = sub_bc.groupby('Model').agg(qty=('Defective Qty(Pair)', 'sum'), cost=('Ttl cost ($)', 'sum')).reset_index()
            mod_data['BC'][s] = [{'model': str(r['Model']).strip(), 'qty': float(r['qty']), 'cost': float(r['cost'])} for _, r in g.iterrows()]
        # RE
        sub_re = sub[sub['_group'] == 'Reinspection']
        if not sub_re.empty:
            g = sub_re.groupby('Model').agg(qty=('Defective Qty(Pair)', 'sum'), hrs=('Working hours', 'sum'), cost=('Ttl cost ($)', 'sum')).reset_index()
            mod_data['RE'][s] = [{'model': str(r['Model']).strip(), 'qty': float(r['qty']), 'hrs': float(r['hrs']), 'cost': float(r['cost'])} for _, r in g.iterrows()]
        # TU
        sub_tu = sub[sub['_group'] == 'Touch-up']
        if not sub_tu.empty:
            g = sub_tu.groupby('Model').agg(qty=('Defective Qty(Pair)', 'sum'), hrs=('Working hours', 'sum'), cost=('Ttl cost ($)', 'sum')).reset_index()
            mod_data['TU'][s] = [{'model': str(r['Model']).strip(), 'qty': float(r['qty']), 'hrs': float(r['hrs']), 'cost': float(r['cost'])} for _, r in g.iterrows()]

    return top_table, blocks, ov_data, fac_data, mod_data


def parse_copq_report(clean_path):
    wb = openpyxl.load_workbook(clean_path, data_only=True)
    if 'COPQ_Report' not in wb.sheetnames:
        raise ValueError(f"Sheet 'COPQ_Report' not found in {clean_path}")
    ws = wb['COPQ_Report']

    # 1. Top table (rows 2 to 9)
    top_table = {}
    for r in range(2, 10):
        fty = ws.cell(r, 1).value
        if fty and str(fty).strip():
            fty_name = str(fty).strip().upper()
            if fty_name == 'JV3':
                fty_name = 'JVB'
            top_table[fty_name] = {
                'tu': ws.cell(r, 2).value,
                're': ws.cell(r, 3).value,
                'bc': ws.cell(r, 4).value,
                'rw': ws.cell(r, 5).value,
                'total': ws.cell(r, 6).value,
                're_tu': ws.cell(r, 7).value,
            }

    # 2. Lower blocks (Rework, Re-Inspection, Touch-up)
    blocks = {}
    current_block = None
    for r in range(11, ws.max_row + 1):
        c1 = ws.cell(r, 1).value
        if c1 in ['Rework', 'Re-Inspection', 'Touch-up']:
            current_block = str(c1)
            blocks[current_block] = {'left': {}, 'right': {}}
            continue
        if current_block:
            v_l_name = ws.cell(r, 1).value
            v_l_val = ws.cell(r, 2).value
            v_r_name = ws.cell(r, 4).value
            v_r_val = ws.cell(r, 5).value
            if v_l_name and str(v_l_name).strip():
                nl = str(v_l_name).strip().upper()
                if nl == 'JV3': nl = 'JVB'
                blocks[current_block]['left'][nl] = v_l_val
            if v_r_name and str(v_r_name).strip():
                nr = str(v_r_name).strip().upper()
                if nr == 'JV3': nr = 'JVB'
                blocks[current_block]['right'][nr] = v_r_val

    return top_table, blocks

def parse_clean_pivots(clean_path):
    wb = openpyxl.load_workbook(clean_path, data_only=True)

    # 1. Overview sheet
    ov_name = None
    for cand in ['COPQ_Overview', 'COPQ_Pivot_Overview']:
        if cand in wb.sheetnames:
            ov_name = cand
            break
    if not ov_name:
        raise ValueError(f"Neither 'COPQ_Overview' nor 'COPQ_Pivot_Overview' found in {clean_path}")
    ws_ov = wb[ov_name]

    ov_data = {}
    for r in range(10, ws_ov.max_row + 1):
        site = ws_ov.cell(r, 2).value
        if site and str(site).strip() and str(site).strip() != 'Grand Total':
            s_name = str(site).strip().upper()
            if s_name == 'JV3': s_name = 'JVB'
            ov_data[s_name] = {
                'bc_q': ws_ov.cell(r, 3).value,
                'bc_c': ws_ov.cell(r, 5).value,
                're_q': ws_ov.cell(r, 9).value,
                're_c': ws_ov.cell(r, 11).value,
                'tu_q': ws_ov.cell(r, 15).value,
                'tu_c': ws_ov.cell(r, 17).value,
            }

    # 2. Top Factory
    if 'COPQ_Pivot_TopFactory' not in wb.sheetnames:
        raise ValueError(f"'COPQ_Pivot_TopFactory' not found in {clean_path}")
    ws_fac = wb['COPQ_Pivot_TopFactory']

    fac_data = {'BC': {}, 'RE': {}, 'TU': {}}
    curr_site = None
    for r in range(10, ws_fac.max_row):
        s = ws_fac.cell(r, 2).value
        if s and str(s).strip():
            curr_site = str(s).strip().upper()
            if curr_site == 'JV3': curr_site = 'JVB'
        fty = ws_fac.cell(r, 3).value
        if fty and str(fty).strip() not in ('', '(blank)'):
            fty_name = str(fty).strip()
            # BC: Qty=4, Cost=6
            bc_q = ws_fac.cell(r, 4).value
            bc_c = ws_fac.cell(r, 6).value
            if bc_c is not None and bc_c > 0:
                fac_data['BC'].setdefault(curr_site, []).append({
                    'factory': fty_name, 'qty': bc_q, 'cost': bc_c
                })
            # RE: Qty=10, Hrs=11, Cost=12
            re_q = ws_fac.cell(r, 10).value
            re_h = ws_fac.cell(r, 11).value
            re_c = ws_fac.cell(r, 12).value
            if re_c is not None and re_c > 0:
                fac_data['RE'].setdefault(curr_site, []).append({
                    'factory': fty_name, 'qty': re_q, 'hrs': re_h, 'cost': re_c
                })
            # TU: Qty=16, Hrs=17, Cost=18
            tu_q = ws_fac.cell(r, 16).value
            tu_h = ws_fac.cell(r, 17).value
            tu_c = ws_fac.cell(r, 18).value
            if tu_c is not None and tu_c > 0:
                fac_data['TU'].setdefault(curr_site, []).append({
                    'factory': fty_name, 'qty': tu_q, 'hrs': tu_h, 'cost': tu_c
                })

    # 3. Top Model & hours lookup from COPQ_Clean
    if 'COPQ_Pivot_TopModel' not in wb.sheetnames:
        raise ValueError(f"'COPQ_Pivot_TopModel' not found in {clean_path}")
    ws_mod = wb['COPQ_Pivot_TopModel']

    # hours lookup from COPQ_Clean
    model_hours_lookup = {}
    if 'COPQ_Clean' in wb.sheetnames:
        df_clean = pd.read_excel(clean_path, sheet_name='COPQ_Clean')
        if 'Site' in df_clean.columns and 'Model' in df_clean.columns and 'Working hours' in df_clean.columns:
            df_clean['Site_M'] = df_clean['Site'].astype(str).str.strip().str.upper().replace({'JV3': 'JVB'})
            df_clean['Model_M'] = df_clean['Model'].astype(str).str.strip()
            df_clean['Type_Group'] = df_clean.get('Type Detail (Grouped)', df_clean.get('Type', '')).astype(str).str.strip()
            # aggregate hours
            g_hrs = df_clean.groupby(['Site_M', 'Model_M', 'Type_Group'])['Working hours'].sum().reset_index()
            for _, r_h in g_hrs.iterrows():
                model_hours_lookup[(r_h['Site_M'], r_h['Model_M'], r_h['Type_Group'])] = r_h['Working hours']

    mod_data = {'BC': {}, 'RE': {}, 'TU': {}}
    curr_site = None
    for r in range(10, ws_mod.max_row):
        s = ws_mod.cell(r, 2).value
        if s and str(s).strip():
            curr_site = str(s).strip().upper()
            if curr_site == 'JV3': curr_site = 'JVB'
        m = ws_mod.cell(r, 3).value
        if m and str(m).strip() not in ('', '(blank)'):
            m_name = str(m).strip()
            # BC: Qty=4, Cost=5
            bc_q = ws_mod.cell(r, 4).value
            bc_c = ws_mod.cell(r, 5).value
            if bc_c is not None and bc_c > 0:
                mod_data['BC'].setdefault(curr_site, []).append({
                    'model': m_name, 'qty': bc_q, 'cost': bc_c
                })
            # RE: Qty=8, Cost=9
            re_q = ws_mod.cell(r, 8).value
            re_c = ws_mod.cell(r, 9).value
            if re_c is not None and re_c > 0:
                re_hrs = model_hours_lookup.get((curr_site, m_name, 'Reinspection'), None)
                mod_data['RE'].setdefault(curr_site, []).append({
                    'model': m_name, 'qty': re_q, 'hrs': re_hrs, 'cost': re_c
                })
            # TU: Qty=12, Cost=13
            tu_q = ws_mod.cell(r, 12).value
            tu_c = ws_mod.cell(r, 13).value
            if tu_c is not None and tu_c > 0:
                tu_hrs = model_hours_lookup.get((curr_site, m_name, 'Touch-up'), None)
                mod_data['TU'].setdefault(curr_site, []).append({
                    'model': m_name, 'qty': tu_q, 'hrs': tu_hrs, 'cost': tu_c
                })

    return ov_data, fac_data, mod_data

def update_copq_database_25(db_path, top_table, blocks, clean_path, report_month):
    wb = openpyxl.load_workbook(db_path)

    month_sheet_name = report_month.strftime("%b-%Y")    # e.g. Aug-2026
    month_title_str = report_month.strftime("%b %Y")     # e.g. Aug 2026

    # 1. Create or replace month sheet by cloning the reference month template
    # (July-2026 or Current) to preserve 100% of the exact cell styles, fonts,
    # borders, fills, formulas, and number formats!
    if month_sheet_name in wb.sheetnames:
        del wb[month_sheet_name]

    # Choose most appropriate reference sheet to clone styles and formulas
    prior_date = (report_month.replace(day=1) - dt.timedelta(days=1))
    ref_candidates = [
        prior_date.strftime("%b-%Y"),
        "Aug-2026",
        "July-2026",
        "Jun-2026",
        "Current"
    ]
    ref_sheet = next((s for s in ref_candidates if s in wb.sheetnames), "Current")
    ws = wb.copy_worksheet(wb[ref_sheet])
    ws.title = month_sheet_name

    # Header
    ws["A1"] = f"COPQ Overview (Based on {month_title_str})"

    sites_order = ["VH", "VH2", "VH3", "VH4", "JV", "JV2", "JVB"]
    fmt_currency = '"$"#,##0.00'
    fmt_qty = '#,##0.0'
    fmt_int = '#,##0'

    # Fill top table rows 3 to 9 (Col B: TU, Col C: RE, Col D: BC, Col E: RW)
    # Col F (=SUM(B:E)) and Col G (=B+C) formulas are already intact from template!
    for i, s in enumerate(sites_order, start=3):
        ws[f"A{i}"] = s
        d = top_table.get(s, {})
        ws[f"B{i}"] = round(d.get("tu", 0), 2) if d.get("tu") else None
        ws[f"C{i}"] = round(d.get("re", 0), 2) if d.get("re") else None
        ws[f"D{i}"] = round(d.get("bc", 0), 2) if d.get("bc") else None
        ws[f"E{i}"] = round(d.get("rw", 0), 2) if d.get("rw") else None
        ws[f"F{i}"] = f"=SUM(B{i}:E{i})"
        ws[f"G{i}"] = f"=B{i}+C{i}"
        ws[f"F{i}"].number_format = fmt_currency
        ws[f"G{i}"].number_format = fmt_currency

    # CLG row 10 formulas are intact from template

    # Side Region table formulas are intact from template

    # Detail blocks
    # 1. Rework / BC Grade Qty (rows 16-19)
    rw_sites = ["JV2", "JV", "VH2", "VH"]
    rw_block = blocks.get("Rework", {})
    for idx, s in enumerate(rw_sites, start=16):
        ws[f"A{idx}"] = s
        ws[f"B{idx}"] = rw_block.get("left", {}).get(s, None)
        ws[f"D{idx}"] = s
        ws[f"E{idx}"] = rw_block.get("right", {}).get(s, None)

    # 2. Re-Inspection (rows 23-28)
    re_sites = ["JVB", "JV2", "JV", "VH3", "VH2", "VH"]
    re_block = blocks.get("Re-Inspection", {})
    for idx, s in enumerate(re_sites, start=23):
        ws[f"A{idx}"] = s
        ws[f"B{idx}"] = re_block.get("left", {}).get(s, None)
        ws[f"D{idx}"] = s
        ws[f"E{idx}"] = re_block.get("right", {}).get(s, None)

    # 3. Touch-up (rows 32-36)
    tu_sites = ["JVB", "JV", "VH4", "VH2", "VH"]
    tu_block = blocks.get("Touch-up", {})
    for idx, s in enumerate(tu_sites, start=32):
        ws[f"A{idx}"] = s
        ws[f"B{idx}"] = tu_block.get("left", {}).get(s, None)
        ws[f"D{idx}"] = s
        ws[f"E{idx}"] = tu_block.get("right", {}).get(s, None)

    # Remove COPQ_rows sheet if present (user requested not to create raw rows sheet)
    rows_sheet_name = f"COPQ_rows_{month_sheet_name}"
    if rows_sheet_name in wb.sheetnames:
        del wb[rows_sheet_name]

    # Refresh sheet 'Current'
    if "Current" in wb.sheetnames:
        del wb["Current"]
    cur = wb.copy_worksheet(ws)
    cur.title = "Current"

    wb.save(db_path)
    return wb

def update_copq_type_analysis(type_path, ov_data, fac_data, mod_data, report_month):
    wb = openpyxl.load_workbook(type_path)
    if "Current month" not in wb.sheetnames:
        raise ValueError(f"Sheet 'Current month' not found in {type_path}")
    ws = wb["Current month"]

    month_hdr = report_month.strftime("%b")             # e.g. Aug
    snapshot_sheet_name = report_month.strftime("%b-%y") # e.g. Aug-26
    month_long_label = report_month.strftime("%b, %Y")   # e.g. Aug, 2026

    # --- 1. Determine column indices by chronological calendar ---
    # In CoPQ_type_analysis: 2026 months start at Col 18 (Jan) through Col 29 (Dec)
    if report_month.year == 2026:
        curr_col = 17 + report_month.month
    else:
        prev_col_scan = None
        for c in range(3, ws.max_column + 1):
            v_hdr = ws.cell(2, c).value
            v_dat = str(ws.cell(27, c).value or "")
            if v_hdr and not v_dat.startswith("="):
                prev_col_scan = c
        curr_col = (prev_col_scan + 1) if prev_col_scan else 25

    prev_col = curr_col - 1
    formula_col = curr_col + 1

    prev_let = get_column_letter(prev_col)
    curr_let = get_column_letter(curr_col)
    form_let = get_column_letter(formula_col)

    # Headers
    for hdr_row in [2, 9, 17, 26, 35, 43]:
        cell = ws.cell(hdr_row, curr_col, month_hdr)
        ref_cell = ws.cell(hdr_row, prev_col)
        cell.font = copy.copy(ref_cell.font)
        cell.alignment = copy.copy(ref_cell.alignment)
        cell.fill = copy.copy(ref_cell.fill)
        cell.border = copy.copy(ref_cell.border)

    # Helper to style and set cell
    def set_cell(row, col, val, num_fmt=None):
        c = ws.cell(row, col)
        c.value = val
        ref = ws.cell(row, prev_col)
        c.font = copy.copy(ref.font)
        c.alignment = copy.copy(ref.alignment)
        c.border = copy.copy(ref.border)
        if num_fmt:
            c.number_format = num_fmt
        elif ref.number_format:
            c.number_format = ref.number_format
        return c

    fmt_curr = '"$"#,##0.00'
    fmt_qty = '#,##0.0'

    # --- BC Grade ---
    # Qty rows 3-6: VH, VH2, JV, JV2
    for r, s in [(3, 'VH'), (4, 'VH2'), (5, 'JV'), (6, 'JV2')]:
        q = ov_data.get(s, {}).get('bc_q')
        set_cell(r, curr_col, q if q else None, fmt_qty)
    # Cost rows 10-13: VH, VH2, JV, JV2
    for r, s in [(10, 'VH'), (11, 'VH2'), (12, 'JV'), (13, 'JV2')]:
        c = ov_data.get(s, {}).get('bc_c')
        set_cell(r, curr_col, round(c, 2) if c else None, fmt_curr)

    # --- Re-inspect ---
    # Qty rows 18-23: VH, VH2, VH3, JV, JV2, JVB
    for r, s in [(18, 'VH'), (19, 'VH2'), (20, 'VH3'), (21, 'JV'), (22, 'JV2'), (23, 'JVB')]:
        q = ov_data.get(s, {}).get('re_q')
        set_cell(r, curr_col, q if q else None, '#,##0')
    # Cost rows 27-32: VH, VH2, VH3, JV, JV2, JVB
    for r, s in [(27, 'VH'), (28, 'VH2'), (29, 'VH3'), (30, 'JV'), (31, 'JV2'), (32, 'JVB')]:
        c = ov_data.get(s, {}).get('re_c')
        set_cell(r, curr_col, round(c, 2) if c else None, fmt_curr)

    # --- Touch-up ---
    # Qty rows 36-40: VH, VH2, VH4, JV, JVB
    for r, s in [(36, 'VH'), (37, 'VH2'), (38, 'VH4'), (39, 'JV'), (40, 'JVB')]:
        q = ov_data.get(s, {}).get('tu_q')
        set_cell(r, curr_col, q if q else None, '#,##0')
    # Cost rows 44-48: VH, VH2, VH4, JV, JVB
    for r, s in [(44, 'VH'), (45, 'VH2'), (46, 'VH4'), (47, 'JV'), (48, 'JVB')]:
        c = ov_data.get(s, {}).get('tu_c')
        set_cell(r, curr_col, round(c, 2) if c else None, fmt_curr)

    # --- 2. Shift difference formulas to formula_col ---
    # Re-inspect rows 27-32
    for r in range(27, 33):
        f_cell = ws.cell(r, formula_col, f"={curr_let}{r}-{prev_let}{r}")
        f_cell.number_format = fmt_curr
        ref = ws.cell(r, curr_col)
        f_cell.font = copy.copy(ref.font)
        f_cell.alignment = copy.copy(ref.alignment)
        f_cell.border = copy.copy(ref.border)
    # Touch-up rows 44-47
    for r in range(44, 48):
        f_cell = ws.cell(r, formula_col, f"={curr_let}{r}-{prev_let}{r}")
        f_cell.number_format = fmt_curr
        ref = ws.cell(r, curr_col)
        f_cell.font = copy.copy(ref.font)
        f_cell.alignment = copy.copy(ref.alignment)
        f_cell.border = copy.copy(ref.border)

    # --- 3. Lower Detail Tables (rows 54 to 105) ---
    # Table 1: BC Factory (rows 55-58) - Sorted DESCENDING by cost
    bc_fac_blocks = [
        ('VH', 2), ('VH2', 8), ('JV', 14), ('JV2', 20)
    ]
    for site, col_start in bc_fac_blocks:
        # clear rows 55-58
        for r in range(55, 59):
            for c_off in range(4):
                ws.cell(r, col_start + c_off).value = None
        # get site items
        items = fac_data['BC'].get(site, [])
        items_sorted = sorted(items, key=lambda x: x['cost'], reverse=True)[:4]
        for idx, it in enumerate(items_sorted):
            r = 55 + idx
            ws.cell(r, col_start).value = site
            ws.cell(r, col_start + 1).value = it['factory']
            c_q = ws.cell(r, col_start + 2)
            c_q.value = it['qty']
            c_q.number_format = fmt_qty
            c_c = ws.cell(r, col_start + 3)
            c_c.value = round(it['cost'], 2)
            c_c.number_format = fmt_curr

    # Table 2: BC Model (rows 62-66) - Top 5 models sorted DESCENDING by cost
    for site, col_start in bc_fac_blocks:
        for r in range(62, 67):
            for c_off in range(4):
                ws.cell(r, col_start + c_off).value = None
        items = mod_data['BC'].get(site, [])
        top5_desc = sorted(items, key=lambda x: x['cost'], reverse=True)[:5]
        for idx, it in enumerate(top5_desc):
            r = 62 + idx
            ws.cell(r, col_start).value = site
            ws.cell(r, col_start + 1).value = it['model']
            c_q = ws.cell(r, col_start + 2)
            c_q.value = it['qty']
            c_q.number_format = fmt_qty
            c_c = ws.cell(r, col_start + 3)
            c_c.value = round(it['cost'], 2)
            c_c.number_format = fmt_curr

    # Table 3: Reinsp Factory (rows 70-76) - Sorted DESCENDING by cost
    re_fac_blocks = [
        ('VH', 2), ('VH2', 8), ('JV', 14), ('JV2', 20), ('JVB', 26), ('VH3', 38)
    ]
    for site, col_start in re_fac_blocks:
        for r in range(70, 77):
            for c_off in range(5):
                ws.cell(r, col_start + c_off).value = None
        items = fac_data['RE'].get(site, [])
        items_sorted = sorted(items, key=lambda x: x['cost'], reverse=True)[:7]
        for idx, it in enumerate(items_sorted):
            r = 70 + idx
            ws.cell(r, col_start).value = site
            ws.cell(r, col_start + 1).value = it['factory']
            c_q = ws.cell(r, col_start + 2)
            c_q.value = it['qty']
            c_q.number_format = '#,##0'
            c_h = ws.cell(r, col_start + 3)
            c_h.value = round(it['hrs'], 1) if it.get('hrs') else None
            c_h.number_format = '0.0'
            c_c = ws.cell(r, col_start + 4)
            c_c.value = round(it['cost'], 2)
            c_c.number_format = fmt_curr

    # Table 4: Reinsp Model (rows 79-83) - Top 5 models sorted DESCENDING by cost
    for site, col_start in re_fac_blocks:
        for r in range(79, 84):
            for c_off in range(5):
                ws.cell(r, col_start + c_off).value = None
        items = mod_data['RE'].get(site, [])
        top5_desc = sorted(items, key=lambda x: x['cost'], reverse=True)[:5]
        for idx, it in enumerate(top5_desc):
            r = 79 + idx
            ws.cell(r, col_start, site)
            ws.cell(r, col_start + 1, it['model'])
            ws.cell(r, col_start + 2, it['qty']).number_format = '#,##0'
            ws.cell(r, col_start + 3, round(it['hrs'], 1) if it.get('hrs') else None).number_format = '0.0'
            ws.cell(r, col_start + 4, round(it['cost'], 2)).number_format = fmt_curr

    # Table 5: Touch-up Factory (rows 88-92) - Sorted DESCENDING by cost
    tu_fac_blocks = [
        ('VH', 2), ('VH2', 8), ('JV', 14), ('JVB', 26), ('VH4', 32)
    ]
    for site, col_start in tu_fac_blocks:
        for r in range(88, 93):
            for c_off in range(5):
                ws.cell(r, col_start + c_off).value = None
        items = fac_data['TU'].get(site, [])
        items_sorted = sorted(items, key=lambda x: x['cost'], reverse=True)[:5]
        for idx, it in enumerate(items_sorted):
            r = 88 + idx
            ws.cell(r, col_start, site)
            ws.cell(r, col_start + 1, it['factory'])
            ws.cell(r, col_start + 2, it['qty']).number_format = '#,##0'
            ws.cell(r, col_start + 3, round(it['hrs'], 1) if it.get('hrs') else None).number_format = '0.0'
            ws.cell(r, col_start + 4, round(it['cost'], 2)).number_format = fmt_curr

    # Table 6: Touch-up Model (rows 96-100) - Models sorted DESCENDING by cost
    for site, col_start in tu_fac_blocks:
        for r in range(96, 101):
            for c_off in range(5):
                ws.cell(r, col_start + c_off).value = None
        items = mod_data['TU'].get(site, [])
        top_desc = sorted(items, key=lambda x: x['cost'], reverse=True)[:5]
        for idx, it in enumerate(top_desc):
            r = 96 + idx
            ws.cell(r, col_start, site)
            ws.cell(r, col_start + 1, it['model'])
            ws.cell(r, col_start + 2, it['qty']).number_format = '#,##0'
            ws.cell(r, col_start + 3, round(it['hrs'], 1) if it.get('hrs') else None).number_format = '0.0'
            ws.cell(r, col_start + 4, round(it['cost'], 2)).number_format = fmt_curr

    # --- 4. Snapshot sheet ---
    if snapshot_sheet_name in wb.sheetnames:
        del wb[snapshot_sheet_name]
    snap = wb.copy_worksheet(ws)
    snap.title = snapshot_sheet_name
    snap["A1"] = f"CoPQ Type Analysis - {month_long_label}"

    wb.save(type_path)
    return wb

import shutil

def run_pipeline(clean_path=None):
    if not clean_path:
        clean_path = CLEAN_XLSX
        if not os.path.exists(clean_path):
            clean_path = os.path.join(BASE, "Output", "COPQ_Clean.xlsx")

    print(f"[+] Using clean source: {clean_path}")
    report_month = detect_report_month(clean_path)
    month_folder = report_month.strftime("%Y-%m")
    print(f"[+] Detected reporting month: {report_month.strftime('%B %Y')} ({month_folder})")

    # 1. Master database accumulates history incrementally without reset

    # 2. Parse source sheets (with fallback to direct parsing from COPQ_Clean)
    try:
        top_table, blocks = parse_copq_report(clean_path)
        ov_data, fac_data, mod_data = parse_clean_pivots(clean_path)
        print(f"[+] Parsed from COPQ_Report and Pivot sheets successfully")
    except Exception as e:
        print(f"[*] Note: Parsing from COPQ_Clean master data directly ({e})")
        top_table, blocks, ov_data, fac_data, mod_data = parse_from_clean_df(clean_path)

    # 3. Update Database workbooks
    print(f"[+] Updating {DB_OVERVIEW}...")
    update_copq_database_25(DB_OVERVIEW, top_table, blocks, clean_path, report_month)

    print(f"[+] Updating {DB_TYPE}...")
    update_copq_type_analysis(DB_TYPE, ov_data, fac_data, mod_data, report_month)

    # 4. Output to Output -> Monthly -> <month_folder>
    out_monthly_dir = os.path.join(BASE, "Output", "Monthly", month_folder)
    out_monthly_db = os.path.join(out_monthly_dir, "Database")
    os.makedirs(out_monthly_db, exist_ok=True)

    shutil.copy2(DB_OVERVIEW, os.path.join(out_monthly_db, os.path.basename(DB_OVERVIEW)))
    shutil.copy2(DB_TYPE, os.path.join(out_monthly_db, os.path.basename(DB_TYPE)))

    # Also copy to root of month folder
    shutil.copy2(DB_OVERVIEW, os.path.join(out_monthly_dir, os.path.basename(DB_OVERVIEW)))
    shutil.copy2(DB_TYPE, os.path.join(out_monthly_dir, os.path.basename(DB_TYPE)))

    print(f"[+] Successfully exported workbooks to:")
    print(f"    - {out_monthly_db}")
    print(f"    - {out_monthly_dir}")
    print(f"    - {os.path.dirname(DB_OVERVIEW)}")

if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else None
    run_pipeline(src)
