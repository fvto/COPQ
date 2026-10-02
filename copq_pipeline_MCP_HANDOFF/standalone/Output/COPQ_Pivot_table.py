"""
COPQ_Pivot_table.py - Script COPQ Pivot Table
=============================================
Python implementation of Output/COPQ_Pivot_table.ts.

Workflow Step 3:
- Targets sheet 'COPQ_Clean'.
- Adds column 'Type Detail (Grouped)': merges Touch-up remark variants
  into 'Touch-up', preserving others from 'Type Detail'.
- Generates 3 structured Pivot views:
    1) COPQ_Pivot_Overview   (Site x Type -> Qty, Hours, Cost)
    2) COPQ_Pivot_TopFactory  (Site + Factory x Type -> Qty, Hours, Cost)
    3) COPQ_Pivot_TopModel    (Site + Model x Type -> Qty, Cost)
  Structured with exact cell coordinates (starting row 10, designated category
  columns) matching write_db_month.py and executive reporting standards.
"""

import os
import sys
import re
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.table import TableColumn
from openpyxl.utils import get_column_letter

# Ensure UTF-8 stdout encoding for Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")

CLEAN_SHEET_NAME = "COPQ_Clean"
F_SITE = "Site"
F_FACTORY = "Factory"
F_MODEL = "Model"
F_TYPE_DETAIL = "Type Detail"
F_REMARK = "Remark (Re-inspection Station)"
F_QTY = "Defective Qty(Pair)"
F_HOURS = "Working hours"
F_COST = "Ttl cost ($)"
F_TYPE_DETAIL_GROUPED = "Type Detail (Grouped)"

TOUCHUP_REMARK_VALUES = [
    "touch-up paint",
    "touch-up paint (bottom from jvb)",
    "touch-up paint (bottom from vh4)"
]

# Standard pivot category columns to preserve layout alignment
STANDARD_CATEGORIES = ["B/C", "Inbound&PI claim", "Reinspection", "Rework", "Touch-up"]


def normalize_text(s: str) -> str:
    """Normalize text: strip zero-width chars, trim, collapse spaces, lower."""
    s = re.sub(r"[\u200B-\u200D\uFEFF\u00A0]", "", str(s or ""))
    s = s.strip()
    s = re.sub(r"\s+", " ", s)
    return s.lower()


def is_touch_up_remark(remark_raw: str) -> bool:
    norm = normalize_text(remark_raw)
    return (
        norm in TOUCHUP_REMARK_VALUES
        or "touch-up" in norm
        or "touch up" in norm
    )


def _to_float(val):
    if val is None or val == "":
        return 0.0
    try:
        s = str(val).replace("$", "").replace(",", "").strip()
        return float(s)
    except (ValueError, TypeError):
        return 0.0


def _style_pivot_sheet(ws, title, is_model=False):
    """Apply professional styling to Pivot sheet headers."""
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    ws["B2"].value = title
    ws["B2"].font = title_font

    filter_lbl_font = Font(name="Calibri", size=9, bold=True, color="595959")
    filter_val_font = Font(name="Calibri", size=9, italic=True)
    ws["B4"].value = "Filter: Remark (Re-inspection Station)"
    ws["B4"].font = filter_lbl_font
    ws["C4"].value = "(All)"
    ws["C4"].font = filter_val_font

    cat_fill = PatternFill(start_color="1F4E78", fill_type="solid")
    cat_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    metric_fill = PatternFill(start_color="2E75B6", fill_type="solid")
    metric_font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")

    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Style row 8 & 9
    for c in range(2, ws.max_column + 1):
        cell_8 = ws.cell(8, c)
        if cell_8.value:
            cell_8.fill = cat_fill
            cell_8.font = cat_font
            cell_8.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_8.border = border_thin

        cell_9 = ws.cell(9, c)
        cell_9.fill = metric_fill
        cell_9.font = metric_font
        cell_9.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_9.border = border_thin

    ws.row_dimensions[8].height = 22
    ws.row_dimensions[9].height = 26
    ws.freeze_panes = "C10" if is_model else "B10"


def run_copq_pivot_table(workbook_or_path):
    """
    Execute Script COPQ Pivot Table.
    Accepts either an openpyxl.Workbook instance or a file path.
    Returns the modified openpyxl.Workbook.
    """
    is_path = isinstance(workbook_or_path, str)
    if is_path:
        wb = openpyxl.load_workbook(workbook_or_path)
    else:
        wb = workbook_or_path

    if CLEAN_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet '{CLEAN_SHEET_NAME}' not found.")

    clean_ws = wb[CLEAN_SHEET_NAME]
    headers = [str(clean_ws.cell(1, c).value or "").strip() for c in range(1, clean_ws.max_column + 1)]

    def col_idx(col_name):
        try:
            return headers.index(col_name) + 1
        except ValueError:
            return -1

    site_c = col_idx(F_SITE)
    factory_c = col_idx(F_FACTORY)
    model_c = col_idx(F_MODEL)
    type_detail_c = col_idx(F_TYPE_DETAIL)
    remark_c = col_idx(F_REMARK)
    qty_c = col_idx(F_QTY)
    hours_c = col_idx(F_HOURS)
    cost_c = col_idx(F_COST)

    if remark_c == -1:
        # Fallback search
        for i, h in enumerate(headers):
            if "remark" in h.lower():
                remark_c = i + 1
                break

    if site_c == -1 or type_detail_c == -1 or remark_c == -1:
        raise ValueError(f"Required columns ('{F_SITE}', '{F_TYPE_DETAIL}', '{F_REMARK}') not found on '{CLEAN_SHEET_NAME}'.")

    # =========================================================================
    # 1. Add/Populate column 'Type Detail (Grouped)'
    # =========================================================================
    if F_TYPE_DETAIL_GROUPED in headers:
        grouped_c = headers.index(F_TYPE_DETAIL_GROUPED) + 1
    else:
        grouped_c = len(headers) + 1
        headers.append(F_TYPE_DETAIL_GROUPED)
        hdr_cell = clean_ws.cell(1, grouped_c, F_TYPE_DETAIL_GROUPED)
        hdr_cell.fill = PatternFill(start_color="1F4E78", fill_type="solid")
        hdr_cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        hdr_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Calibri", size=9)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    clean_records = []
    for r in range(2, clean_ws.max_row + 1):
        site_val = str(clean_ws.cell(r, site_c).value or "").strip()
        factory_val = str(clean_ws.cell(r, factory_c).value or "").strip() if factory_c > 0 else ""
        model_val = str(clean_ws.cell(r, model_c).value or "").strip() if model_c > 0 else ""
        type_detail_val = str(clean_ws.cell(r, type_detail_c).value or "").strip()
        remark_val = str(clean_ws.cell(r, remark_c).value or "").strip()
        qty_val = _to_float(clean_ws.cell(r, qty_c).value) if qty_c > 0 else 0.0
        hours_val = _to_float(clean_ws.cell(r, hours_c).value) if hours_c > 0 else 0.0
        cost_val = _to_float(clean_ws.cell(r, cost_c).value) if cost_c > 0 else 0.0

        grouped_val = "Touch-up" if is_touch_up_remark(remark_val) else type_detail_val

        cell = clean_ws.cell(r, grouped_c, grouped_val)
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="left", vertical="center")

        clean_records.append({
            "site": site_val,
            "factory": factory_val,
            "model": model_val,
            "grouped": grouped_val,
            "qty": qty_val,
            "hours": hours_val,
            "cost": cost_val,
        })

    # Update Table
    for t_name in list(clean_ws.tables.keys()):
        tab = clean_ws.tables[t_name]
        new_col_letter = get_column_letter(max(grouped_c, clean_ws.max_column))
        tab.ref = f"A1:{new_col_letter}{clean_ws.max_row}"
        existing_col_names = [col.name for col in tab.tableColumns]
        if F_TYPE_DETAIL_GROUPED not in existing_col_names:
            tab.tableColumns.append(TableColumn(id=len(tab.tableColumns) + 1, name=F_TYPE_DETAIL_GROUPED))

    clean_ws.column_dimensions[get_column_letter(grouped_c)].width = 20

    # Determine all categories present, maintaining exact standard order
    categories = list(STANDARD_CATEGORIES)
    data_cats = {r["grouped"] for r in clean_records if r["grouped"]}
    for c in sorted(data_cats):
        if c not in categories:
            categories.append(c)

    # Reusable formatters
    border_data = thin_border
    total_border = Border(
        top=Side(style="thin", color="000000"),
        bottom=Side(style="double", color="000000"),
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
    )
    bold_font = Font(name="Calibri", size=9, bold=True)
    num_fmt_qty = "#,##0"
    num_fmt_hrs = "#,##0.00"
    num_fmt_cost = "$#,##0.00"

    # =========================================================================
    # 2. Sheet 1: COPQ_Pivot_Overview
    # =========================================================================
    ov_name = "COPQ_Pivot_Overview"
    if ov_name in wb.sheetnames:
        del wb[ov_name]
    ws_ov = wb.create_sheet(title=ov_name)

    # Set headers
    ws_ov.cell(8, 2, "Site")
    ws_ov.cell(9, 2, "Row Labels")

    col_ptr = 3
    cat_col_map_ov = {}
    for cat in categories:
        cat_col_map_ov[cat] = col_ptr
        ws_ov.cell(8, col_ptr, cat)
        ws_ov.cell(9, col_ptr, "Sum of Defective Qty(Pair)")
        ws_ov.cell(9, col_ptr + 1, "Sum of Working hours")
        ws_ov.cell(9, col_ptr + 2, "Sum of Ttl cost ($)")
        col_ptr += 3

    # Total columns
    total_col_ov = col_ptr
    ws_ov.cell(8, total_col_ov, "Total")
    ws_ov.cell(9, total_col_ov, "Sum of Defective Qty(Pair)")
    ws_ov.cell(9, total_col_ov + 1, "Sum of Working hours")
    ws_ov.cell(9, total_col_ov + 2, "Sum of Ttl cost ($)")

    # Aggregate Overview
    site_order = ["JV", "JV2", "JV3", "VH", "VH2", "VH3", "VH4"]
    present_sites = {r["site"] for r in clean_records if r["site"]}
    ordered_sites = [s for s in site_order if s in present_sites] + sorted([s for s in present_sites if s not in site_order])

    curr_row = 10
    col_totals_ov = [0.0] * (total_col_ov + 3)

    for s in ordered_sites:
        ws_ov.cell(curr_row, 2, s).font = bold_font
        ws_ov.cell(curr_row, 2).border = border_data
        row_qty, row_hrs, row_cost = 0.0, 0.0, 0.0

        for cat in categories:
            c_idx = cat_col_map_ov[cat]
            sub = [r for r in clean_records if r["site"] == s and r["grouped"] == cat]
            q = sum(r["qty"] for r in sub)
            h = sum(r["hours"] for r in sub)
            c = sum(r["cost"] for r in sub)

            cell_q = ws_ov.cell(curr_row, c_idx, q if q > 0 else "")
            cell_h = ws_ov.cell(curr_row, c_idx + 1, h if h > 0 else "")
            cell_c = ws_ov.cell(curr_row, c_idx + 2, c if c > 0 else "")

            cell_q.number_format = num_fmt_qty
            cell_h.number_format = num_fmt_hrs
            cell_c.number_format = num_fmt_cost

            for cell in (cell_q, cell_h, cell_c):
                cell.font = data_font
                cell.border = border_data

            col_totals_ov[c_idx] += q
            col_totals_ov[c_idx + 1] += h
            col_totals_ov[c_idx + 2] += c

            row_qty += q
            row_hrs += h
            row_cost += c

        # Row totals
        cell_tq = ws_ov.cell(curr_row, total_col_ov, row_qty if row_qty > 0 else "")
        cell_th = ws_ov.cell(curr_row, total_col_ov + 1, row_hrs if row_hrs > 0 else "")
        cell_tc = ws_ov.cell(curr_row, total_col_ov + 2, row_cost if row_cost > 0 else "")

        cell_tq.number_format = num_fmt_qty
        cell_th.number_format = num_fmt_hrs
        cell_tc.number_format = num_fmt_cost

        for cell in (cell_tq, cell_th, cell_tc):
            cell.font = bold_font
            cell.border = border_data

        col_totals_ov[total_col_ov] += row_qty
        col_totals_ov[total_col_ov + 1] += row_hrs
        col_totals_ov[total_col_ov + 2] += row_cost

        curr_row += 1

    # Grand Total row
    ws_ov.cell(curr_row, 2, "Grand Total").font = bold_font
    ws_ov.cell(curr_row, 2).border = total_border
    for c in range(3, total_col_ov + 3):
        val = col_totals_ov[c]
        cell = ws_ov.cell(curr_row, c, val if val > 0 else "")
        cell.font = bold_font
        cell.border = total_border
        if (c - 3) % 3 == 0:
            cell.number_format = num_fmt_qty
        elif (c - 3) % 3 == 1:
            cell.number_format = num_fmt_hrs
        else:
            cell.number_format = num_fmt_cost

    _style_pivot_sheet(ws_ov, "Overview - by Site & Type")

    # =========================================================================
    # 3. Sheet 2: COPQ_Pivot_TopFactory
    # =========================================================================
    fac_name = "COPQ_Pivot_TopFactory"
    if fac_name in wb.sheetnames:
        del wb[fac_name]
    ws_fac = wb.create_sheet(title=fac_name)

    ws_fac.cell(8, 2, "Site")
    ws_fac.cell(8, 3, "Factory")
    ws_fac.cell(9, 2, "Site")
    ws_fac.cell(9, 3, "Factory")

    col_ptr = 4
    cat_col_map_fac = {}
    for cat in categories:
        cat_col_map_fac[cat] = col_ptr
        ws_fac.cell(8, col_ptr, cat)
        ws_fac.cell(9, col_ptr, "Sum of Defective Qty(Pair)")
        ws_fac.cell(9, col_ptr + 1, "Sum of Working hours")
        ws_fac.cell(9, col_ptr + 2, "Sum of Ttl cost ($)")
        col_ptr += 3

    total_col_fac = col_ptr
    ws_fac.cell(8, total_col_fac, "Total")
    ws_fac.cell(9, total_col_fac, "Sum of Defective Qty(Pair)")
    ws_fac.cell(9, total_col_fac + 1, "Sum of Working hours")
    ws_fac.cell(9, total_col_fac + 2, "Sum of Ttl cost ($)")

    # Aggregate by (Site, Factory)
    pairs = sorted(list({(r["site"], r["factory"]) for r in clean_records if r["site"] and r["factory"]}))
    curr_row = 10
    col_totals_fac = [0.0] * (total_col_fac + 3)

    for (s, fty) in pairs:
        ws_fac.cell(curr_row, 2, s).font = bold_font
        ws_fac.cell(curr_row, 2).border = border_data
        ws_fac.cell(curr_row, 3, fty).font = data_font
        ws_fac.cell(curr_row, 3).border = border_data

        row_qty, row_hrs, row_cost = 0.0, 0.0, 0.0

        for cat in categories:
            c_idx = cat_col_map_fac[cat]
            sub = [r for r in clean_records if r["site"] == s and r["factory"] == fty and r["grouped"] == cat]
            q = sum(r["qty"] for r in sub)
            h = sum(r["hours"] for r in sub)
            c = sum(r["cost"] for r in sub)

            cell_q = ws_fac.cell(curr_row, c_idx, q if q > 0 else "")
            cell_h = ws_fac.cell(curr_row, c_idx + 1, h if h > 0 else "")
            cell_c = ws_fac.cell(curr_row, c_idx + 2, c if c > 0 else "")

            cell_q.number_format = num_fmt_qty
            cell_h.number_format = num_fmt_hrs
            cell_c.number_format = num_fmt_cost

            for cell in (cell_q, cell_h, cell_c):
                cell.font = data_font
                cell.border = border_data

            col_totals_fac[c_idx] += q
            col_totals_fac[c_idx + 1] += h
            col_totals_fac[c_idx + 2] += c

            row_qty += q
            row_hrs += h
            row_cost += c

        cell_tq = ws_fac.cell(curr_row, total_col_fac, row_qty if row_qty > 0 else "")
        cell_th = ws_fac.cell(curr_row, total_col_fac + 1, row_hrs if row_hrs > 0 else "")
        cell_tc = ws_fac.cell(curr_row, total_col_fac + 2, row_cost if row_cost > 0 else "")

        cell_tq.number_format = num_fmt_qty
        cell_th.number_format = num_fmt_hrs
        cell_tc.number_format = num_fmt_cost

        for cell in (cell_tq, cell_th, cell_tc):
            cell.font = bold_font
            cell.border = border_data

        col_totals_fac[total_col_fac] += row_qty
        col_totals_fac[total_col_fac + 1] += row_hrs
        col_totals_fac[total_col_fac + 2] += row_cost

        curr_row += 1

    # Grand total
    ws_fac.cell(curr_row, 2, "Grand Total").font = bold_font
    ws_fac.cell(curr_row, 2).border = total_border
    ws_fac.cell(curr_row, 3, "").border = total_border
    for c in range(4, total_col_fac + 3):
        val = col_totals_fac[c]
        cell = ws_fac.cell(curr_row, c, val if val > 0 else "")
        cell.font = bold_font
        cell.border = total_border
        if (c - 4) % 3 == 0:
            cell.number_format = num_fmt_qty
        elif (c - 4) % 3 == 1:
            cell.number_format = num_fmt_hrs
        else:
            cell.number_format = num_fmt_cost

    _style_pivot_sheet(ws_fac, "Top Factory - by Site, Factory & Type")

    # =========================================================================
    # 4. Sheet 3: COPQ_Pivot_TopModel
    # =========================================================================
    mod_name = "COPQ_Pivot_TopModel"
    if mod_name in wb.sheetnames:
        del wb[mod_name]
    ws_mod = wb.create_sheet(title=mod_name)

    ws_mod.cell(8, 2, "Site")
    ws_mod.cell(8, 3, "Model")
    ws_mod.cell(9, 2, "Site")
    ws_mod.cell(9, 3, "Model")

    col_ptr = 4
    cat_col_map_mod = {}
    for cat in categories:
        cat_col_map_mod[cat] = col_ptr
        ws_mod.cell(8, col_ptr, cat)
        ws_mod.cell(9, col_ptr, "Sum of Defective Qty(Pair)")
        ws_mod.cell(9, col_ptr + 1, "Sum of Ttl cost ($)")
        col_ptr += 2

    total_col_mod = col_ptr
    ws_mod.cell(8, total_col_mod, "Total")
    ws_mod.cell(9, total_col_mod, "Sum of Defective Qty(Pair)")
    ws_mod.cell(9, total_col_mod + 1, "Sum of Ttl cost ($)")

    # Aggregate by (Site, Model)
    model_pairs = sorted(list({(r["site"], r["model"]) for r in clean_records if r["site"] and r["model"]}))
    curr_row = 10
    col_totals_mod = [0.0] * (total_col_mod + 2)

    for (s, m) in model_pairs:
        ws_mod.cell(curr_row, 2, s).font = bold_font
        ws_mod.cell(curr_row, 2).border = border_data
        ws_mod.cell(curr_row, 3, m).font = data_font
        ws_mod.cell(curr_row, 3).border = border_data

        row_qty, row_cost = 0.0, 0.0

        for cat in categories:
            c_idx = cat_col_map_mod[cat]
            sub = [r for r in clean_records if r["site"] == s and r["model"] == m and r["grouped"] == cat]
            q = sum(r["qty"] for r in sub)
            c = sum(r["cost"] for r in sub)

            cell_q = ws_mod.cell(curr_row, c_idx, q if q > 0 else "")
            cell_c = ws_mod.cell(curr_row, c_idx + 1, c if c > 0 else "")

            cell_q.number_format = num_fmt_qty
            cell_c.number_format = num_fmt_cost

            for cell in (cell_q, cell_c):
                cell.font = data_font
                cell.border = border_data

            col_totals_mod[c_idx] += q
            col_totals_mod[c_idx + 1] += c

            row_qty += q
            row_cost += c

        cell_tq = ws_mod.cell(curr_row, total_col_mod, row_qty if row_qty > 0 else "")
        cell_tc = ws_mod.cell(curr_row, total_col_mod + 1, row_cost if row_cost > 0 else "")

        cell_tq.number_format = num_fmt_qty
        cell_tc.number_format = num_fmt_cost

        for cell in (cell_tq, cell_tc):
            cell.font = bold_font
            cell.border = border_data

        col_totals_mod[total_col_mod] += row_qty
        col_totals_mod[total_col_mod + 1] += row_cost

        curr_row += 1

    # Grand total
    ws_mod.cell(curr_row, 2, "Grand Total").font = bold_font
    ws_mod.cell(curr_row, 2).border = total_border
    ws_mod.cell(curr_row, 3, "").border = total_border
    for c in range(4, total_col_mod + 2):
        val = col_totals_mod[c]
        cell = ws_mod.cell(curr_row, c, val if val > 0 else "")
        cell.font = bold_font
        cell.border = total_border
        if (c - 4) % 2 == 0:
            cell.number_format = num_fmt_qty
        else:
            cell.number_format = num_fmt_cost

    _style_pivot_sheet(ws_mod, "Top Model - by Site, Model & Type", is_model=True)

    # Auto-fit columns across all 3 sheets
    for ws_p in (ws_ov, ws_fac, ws_mod):
        for col in ws_p.columns:
            max_len = 10
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, min(len(str(cell.value)), 35))
            col_letter = get_column_letter(col[0].column)
            ws_p.column_dimensions[col_letter].width = max(max_len + 2, 12)

    print(f"[+] Script COPQ Pivot Table completed: Created {ov_name}, {fac_name}, {mod_name}.")

    if is_path:
        wb.save(workbook_or_path)
    return wb


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), "COPQ_Clean.xlsx")
    if not os.path.exists(target):
        print(f"Target file not found: {target}")
        sys.exit(1)
    run_copq_pivot_table(target)
