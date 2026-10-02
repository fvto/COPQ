"""
Combined_COPQ.py - Script COPQ Clean
====================================
Python implementation of Output/Combined_COPQ.ts.

Workflow Step 1:
- Gathers data across all individual site sheets (excluding COPQ_Clean).
- Builds unified master header list.
- Validates data rows against quality and SOP rules.
- Generates COPQ_Clean master sheet with Table formatting, status columns,
  and highlights validation errors in soft red (#FFC7CE / #9C0006).
"""

import os
import sys
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

# Ensure UTF-8 stdout encoding for Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")

HOURS_FIELD = "Working hours"
QTY_FIELD = "Defective Qty(Pair)"
COST_FIELD = "Ttl cost ($)"
TYPE_FIELD = "Type"
CATEGORY_FIELD = "Category"
STYLE_FIELD = "Style Nbr"
MODEL_FIELD = "Model"

HOURS_LIMIT = 15   # error if hours > 15
QTY_LIMIT = 2000   # error if qty < 2000 (combined with hours > 15)
COST_LIMIT = 1000  # error if cost > 1000
REINSPECTION_TYPE = "reinspection"

REQUIRED_FIELDS = ["Date", "Category", "Type", "Factory", "Model"]
STYLE_FACTORIES = ["factory 1", "factory 2", "factory 3", "factory 4", "factory 5", "sandals"]
MOLD_FACTORIES = ["ip", "os", "stockfitting"]


def _to_float(val):
    if val is None or val == "":
        return None
    try:
        s = str(val).replace("$", "").replace(",", "").strip()
        return float(s)
    except (ValueError, TypeError):
        return None


def run_copq_clean(workbook_or_path):
    """
    Execute Script COPQ Clean.
    Accepts either an openpyxl.Workbook instance or a file path.
    Returns the modified openpyxl.Workbook.
    """
    is_path = isinstance(workbook_or_path, str)
    if is_path:
        wb = openpyxl.load_workbook(workbook_or_path)
    else:
        wb = workbook_or_path

    output_name = "COPQ_Clean"
    skip_sheets = [output_name]

    # Delete existing sheet if present
    if output_name in wb.sheetnames:
        del wb[output_name]

    out_ws = wb.create_sheet(title=output_name)

    # 1. Pass 1: Gather all unique headers across all site sheets
    master_headers = []
    seen_headers = set()

    for sheet_name in wb.sheetnames:
        if sheet_name in skip_sheets:
            continue
        ws = wb[sheet_name]
        if ws.max_row < 3:
            continue

        # Row 2 contains column headers (Row 1 is Title "Cost Of Poor Quality")
        row_vals = [str(ws.cell(2, c).value or "").strip() for c in range(1, ws.max_column + 1)]
        for h in row_vals:
            if h and h not in seen_headers:
                seen_headers.add(h)
                master_headers.append(h)

    if not master_headers:
        print("[!] No valid site data found across sheets.")
        if is_path:
            wb.save(workbook_or_path)
        return wb

    full_headers = ["Site"] + master_headers + ["Missing Fields", "Validation Errors", "Data Quality Status"]
    combined_rows = []
    highlight_cells = []  # list of (row_idx_1_based, col_idx_1_based)

    # 2. Pass 2: Map row data dynamically into the uniform master header structure
    for sheet_name in wb.sheetnames:
        if sheet_name in skip_sheets:
            continue
        ws = wb[sheet_name]
        if ws.max_row < 3:
            continue

        raw_sheet_headers = [str(ws.cell(2, c).value or "").strip() for c in range(1, ws.max_column + 1)]

        def get_col_idx(field_name):
            try:
                return raw_sheet_headers.index(field_name)
            except ValueError:
                return -1

        hours_idx = get_col_idx(HOURS_FIELD)
        qty_idx = get_col_idx(QTY_FIELD)
        cost_idx = get_col_idx(COST_FIELD)
        type_idx = get_col_idx(TYPE_FIELD)
        category_idx = get_col_idx(CATEGORY_FIELD)
        style_idx = get_col_idx(STYLE_FIELD)
        model_idx = get_col_idx(MODEL_FIELD)
        factory_idx = get_col_idx("Factory")
        mold_idx = get_col_idx("Mold")
        manpower_idx = get_col_idx("Manpower")

        for r in range(3, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if all(v is None or str(v).strip() == "" for v in row_vals):
                continue

            row_map = {}
            for i, h in enumerate(raw_sheet_headers):
                if h and i < len(row_vals):
                    row_map[h] = row_vals[i]

            # Exclude Rework records per user requirement
            type_raw = str(row_map.get(TYPE_FIELD, "") or "").strip().lower()
            if type_raw == "rework":
                continue

            normalized = [row_map.get(h, "") for h in master_headers]

            missing = [
                h for h in REQUIRED_FIELDS
                if h not in row_map or row_map[h] is None or str(row_map[h]).strip() == ""
            ]

            # --- Validation checks ---
            validation_errors = []
            local_highlights = []

            def get_val(idx):
                return row_vals[idx] if 0 <= idx < len(row_vals) else None

            hours_val = _to_float(get_val(hours_idx))
            qty_val = _to_float(get_val(qty_idx))
            cost_val = _to_float(get_val(cost_idx))
            manpower_val = _to_float(get_val(manpower_idx))

            has_hours = hours_val is not None
            has_qty = qty_val is not None
            has_cost = cost_val is not None

            # Rule 1: Working hours > 15 AND Defective Qty(Pair) < 2000
            if has_hours and has_qty and hours_val > HOURS_LIMIT and qty_val < QTY_LIMIT:
                validation_errors.append(f"{HOURS_FIELD} > {HOURS_LIMIT} & {QTY_FIELD} < {QTY_LIMIT}")
                local_highlights.extend([HOURS_FIELD, QTY_FIELD])

            # Rule 2: Ttl cost ($) > 1000
            if has_cost and cost_val > COST_LIMIT:
                validation_errors.append(f"{COST_FIELD} > {COST_LIMIT}")
                local_highlights.append(COST_FIELD)

            type_raw = str(get_val(type_idx) or "").strip().lower()
            is_reinspection = (type_raw == REINSPECTION_TYPE)

            factory_raw = str(get_val(factory_idx) or "").strip().lower()
            mold_raw = get_val(mold_idx)
            mold_blank = (mold_idx >= 0 and (mold_raw is None or str(mold_raw).strip() == ""))

            uses_style_nbr = factory_raw in STYLE_FACTORIES
            uses_mold = factory_raw in MOLD_FACTORIES

            category_raw = get_val(category_idx)
            category_blank = (category_idx >= 0 and (category_raw is None or str(category_raw).strip() == ""))

            type_blank = (type_idx >= 0 and (type_raw == ""))

            style_raw = get_val(style_idx)
            style_blank = (style_idx >= 0 and (style_raw is None or str(style_raw).strip() == ""))

            model_raw = get_val(model_idx)
            model_blank = (model_idx >= 0 and (model_raw is None or str(model_raw).strip() == ""))

            hours_blank_or_zero = (hours_idx >= 0 and (not has_hours or hours_val == 0))
            qty_blank_or_zero = (qty_idx >= 0 and (not has_qty or qty_val == 0))
            cost_blank_or_zero = (cost_idx >= 0 and (not has_cost or cost_val == 0))

            manpower_raw = get_val(manpower_idx)
            manpower_blank_or_zero = (
                manpower_idx >= 0 and (manpower_raw is None or str(manpower_raw).strip() == "" or manpower_val == 0)
            )

            identifier_blank = mold_blank if uses_mold else style_blank
            identifier_idx = mold_idx if uses_mold else style_idx

            present_fields = [i for i in [category_idx, identifier_idx, model_idx, manpower_idx, hours_idx, cost_idx] if i >= 0]
            blank_flags = [category_blank, identifier_blank, model_blank, manpower_blank_or_zero, hours_blank_or_zero, cost_blank_or_zero]
            blank_count = sum(1 for (i, b) in zip([category_idx, identifier_idx, model_idx, manpower_idx, hours_idx, cost_idx], blank_flags) if i >= 0 and b)

            is_blank_data_row = len(present_fields) >= 4 and blank_count == len(present_fields)

            if is_blank_data_row:
                validation_errors.append("Blank data")
                local_highlights.extend([
                    CATEGORY_FIELD,
                    "Mold" if uses_mold else STYLE_FIELD,
                    MODEL_FIELD,
                    "Manpower",
                    HOURS_FIELD,
                    COST_FIELD
                ])
            else:
                if is_reinspection and hours_blank_or_zero:
                    validation_errors.append(f"Wrong {HOURS_FIELD} => Wrong {COST_FIELD}")
                    local_highlights.extend([HOURS_FIELD, COST_FIELD])

                if is_reinspection and qty_blank_or_zero:
                    validation_errors.append(f"Wrong {QTY_FIELD}")
                    local_highlights.append(QTY_FIELD)

                if is_reinspection and uses_style_nbr and style_blank:
                    validation_errors.append(f"{STYLE_FIELD} blank (Reinspection)")
                    local_highlights.append(STYLE_FIELD)

                if is_reinspection and uses_mold and mold_blank:
                    validation_errors.append("Mold blank (Reinspection)")
                    local_highlights.append("Mold")

                if is_reinspection and model_blank:
                    validation_errors.append(f"{MODEL_FIELD} blank (Reinspection)")
                    local_highlights.append(MODEL_FIELD)

                if type_blank and cost_blank_or_zero:
                    validation_errors.append(f"Blank {TYPE_FIELD} => Blank {COST_FIELD}")
                    local_highlights.extend([TYPE_FIELD, COST_FIELD])

                if category_blank:
                    validation_errors.append(f"{CATEGORY_FIELD} blank")
                    local_highlights.append(CATEGORY_FIELD)

                if type_blank and not cost_blank_or_zero:
                    validation_errors.append(f"{TYPE_FIELD} blank")
                    local_highlights.append(TYPE_FIELD)

            # Data Quality Status
            if missing and validation_errors:
                status = "Missing Data & Validation Error"
            elif missing:
                status = "Missing Data"
            elif validation_errors:
                status = "Validation Error"
            else:
                status = "OK"

            combined_rows.append([
                sheet_name,
                *normalized,
                ", ".join(missing),
                ", ".join(validation_errors),
                status
            ])

            row_idx_in_combined = len(combined_rows) + 1  # 1-based, header is row 1
            for field in set(local_highlights):
                if field in master_headers:
                    col_idx_in_combined = 2 + master_headers.index(field)  # Col 1 is Site
                    highlight_cells.append((row_idx_in_combined, col_idx_in_combined))

    if not combined_rows:
        print("[!] No data rows to combine.")
        if is_path:
            wb.save(workbook_or_path)
        return wb

    # Write headers and rows to COPQ_Clean
    out_ws.append(full_headers)
    for row in combined_rows:
        out_ws.append(row)

    # Style Header
    hdr_fill = PatternFill(start_color="1F4E78", fill_type="solid")
    hdr_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for c in range(1, len(full_headers) + 1):
        cell = out_ws.cell(1, c)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Add Table
    max_row = len(combined_rows) + 1
    max_col_letter = get_column_letter(len(full_headers))
    table = Table(displayName="COPQ_CleanData", ref=f"A1:{max_col_letter}{max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )
    out_ws.add_table(table)

    # Apply highlight styling to failing cells
    fail_fill = PatternFill(start_color="FFC7CE", fill_type="solid")
    fail_font = Font(name="Calibri", size=9, color="9C0006")

    for r_idx, c_idx in highlight_cells:
        cell = out_ws.cell(r_idx, c_idx)
        cell.fill = fail_fill
        cell.font = fail_font

    # Freeze panes below header
    out_ws.freeze_panes = "A2"

    # Auto-fit column widths
    for col in out_ws.columns:
        max_len = 10
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)), 45))
        col_letter = get_column_letter(col[0].column)
        out_ws.column_dimensions[col_letter].width = max_len + 3

    print(f"[+] Script COPQ Clean completed: {len(combined_rows)} rows combined on 'COPQ_Clean' with {len(highlight_cells)} highlighted issues.")

    if is_path:
        wb.save(workbook_or_path)
    return wb


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), "COPQ_Clean.xlsx")
    if not os.path.exists(target):
        print(f"Target file not found: {target}")
        sys.exit(1)
    run_copq_clean(target)
