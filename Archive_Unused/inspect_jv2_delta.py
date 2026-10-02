 #!/usr/bin/env python3
"""Inspect the JV2 delta text on the deck template."""

import os

from pptx import Presentation

ppts = [
    os.path.join("Template_COPQ", "Template_COPQ.pptx"),
    os.path.join("..", "Jul,2026 CoPQ cost trending.pptx"),
]

ppt_path = next((p for p in ppts if os.path.exists(p)), ppts[0])
prs = Presentation(ppt_path)
slide2 = prs.slides[2]  # Slide 3 in the deck

for shape in slide2.shapes:
    if shape.name == "Group 92":
        rects = [s for s in shape.shapes if "Rectangle" in s.name]
        if rects:
            text = rects[0].text_frame.text
            print(f"JV2 delta text in {ppt_path}: {text!r}")
            break
else:
    print(f"No 'Group 92' shape found in {ppt_path}.")
