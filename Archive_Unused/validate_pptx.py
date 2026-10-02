#!/usr/bin/env python3
"""Validate the generated COPQ_Report_Aug_2026.pptx: structure preserved + data updated."""
import os
import sys
from pptx import Presentation
from pptx.oxml.ns import qn
import openpyxl, pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Output", "COPQ_Report_Aug_2026.pptx")
SRC = os.path.join(HERE, "Template_COPQ", "Template_COPQ.pptx")

def chart_summary(prs, sidx):
    out = []
    for shp in prs.slides[sidx].shapes:
        if shp.has_chart:
            ch = shp.chart
            title = None
            try:
                if ch.has_title and ch.chart_title: title = ch.chart_title.text_frame.text
            except Exception: pass
            sers = []
            for ser in ch._chartSpace.chart.plotArea.iter(qn('c:ser')):
                tx = ser.find(qn('c:tx'))
                nm = None
                if tx is not None:
                    sr = tx.find(qn('c:strRef'))
                    if sr is not None:
                        c = sr.find(qn('c:strCache'))
                        if c is not None and c.find(qn('c:pt')) is not None:
                            v = c.find(qn('c:pt')).find(qn('c:v'))
                            nm = v.text if v is not None else None
                val = ser.find(qn('c:val'))
                n = 0
                if val is not None:
                    nr = val.find(qn('c:numRef'))
                    if nr is not None and nr.find(qn('c:numCache')) is not None:
                        n = len(nr.find(qn('c:numCache')).findall(qn('c:pt')))
                sers.append((nm, n))
            out.append((shp.name, str(ch.chart_type), title, sers))
    return out

src = Presentation(SRC)
out = Presentation(OUT)
print(f"SLIDES: src={len(src.slides)} out={len(out.slides)}  -> {'OK' if len(src.slides)==len(out.slides) else 'FAIL'}")

# title
t0 = out.slides[0].shapes[27].text_frame.text if False else None
for shp in out.slides[0].shapes:
    if shp.has_text_frame and "Cost of Poor Quality" in shp.text_frame.text:
        print("TITLE text:", shp.text_frame.text.replace("\n"," | "))

# table totals
for shp in out.slides[0].shapes:
    if shp.has_table:
        row0 = [c.text.replace("\n"," ") for c in shp.table.rows[0].cells]
        print("TABLE R0:", row0)

# chart type preservation + series counts per slide
for i in (0,1):
    sc = chart_summary(src, i)
    oc = chart_summary(out, i)
    smap = {n:(t,s) for n,t,ti,s in sc}
    print(f"\n--- SLIDE {i} chart types/series (src vs out) ---")
    ok = True
    for name,t,ti,sers in oc:
        st, ss = smap.get(name, ("?","?"))
        same_type = (st==t)
        # compare series count
        ssrc = dict((nm,n) for nm,n in ss)
        cnt_ok = all(ssrc.get(nm,n)==n for nm,n in sers)
        if not same_type: ok=False
        print(f"  {name:10} type={t:14} {'=' if same_type else 'CHANGED!'} series={sers}")
    print(f"  chart-type-preserved: {'OK' if ok else 'FAIL'}")

# verify defect colors written
print("\n--- Top-3 Defect series colors (slide1) ---")
for shp in out.slides[1].shapes:
    if shp.has_chart and shp.name in ("Chart 78","Chart 80","Chart 92","Chart 96"):
        print(f" {shp.name}:")
        for ser in shp.chart._chartSpace.chart.plotArea.iter(qn('c:ser')):
            tx = ser.find(qn('c:tx'))
            nm=None
            if tx is not None:
                sr=tx.find(qn('c:strRef'))
                if sr is not None:
                    c=sr.find(qn('c:strCache'))
                    if c is not None and c.find(qn('c:pt')) is not None:
                        v=c.find(qn('c:pt')).find(qn('c:v')); nm=v.text if v is not None else None
            col=None
            sp=ser.find(qn('c:spPr'))
            if sp is not None:
                sf=sp.find(qn('a:solidFill'))
                if sf is not None:
                    srgb=sf.find(qn('a:srgbClr'))
                    if srgb is not None: col=srgb.get('val')
            print(f"     {nm} -> #{col}")
print("\nVALIDATION DONE")
