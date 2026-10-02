#!/usr/bin/env python3
"""
build_report_workbook.py
========================
Build the SINGLE LIVE-SOURCE Excel workbook that the PowerPoint report links to.
One sheet per slide's data. Fed from:
  * Database/CoPQ database 25.xlsx      (overview + row-level COPQ for the month)
  * Database/CoPQ_type_analysis.xlsx   (rolling monthly trend)
  * Output/FTT_Combined_Report.xlsx     (FTT defect charts -- not in the 2 DB files)

Output: Output/COPQ_Report_Data.xlsx  (the workbook the .pptx LINKS to)

Sheets (one per slide of the deck):
  Slide1_Overview     -> site totals (TU/RE/BC/excl-Rework) + VN/ID/CLG
  Slide2_BCGrade      -> BC Grade cost by model, per site (from row data)
  Slide3_Reinspect    -> Re-inspection cost by model, per site
  Slide3_Touchup      -> Touch-up paint cost by model, per site
  Slide4_Monthly      -> rolling monthly TU/RE/BC per site (from type_analysis)
  Slide4_TopDefect    -> FTT Top Defects Summary (from FTT_Combined_Report)
  Slide4_Top3Model    -> Top-3 models by cost per site (from row data)

Configurable at the top. Idempotent: re-running overwrites the sheets.
"""
import os
import sys
import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(HERE)
if BASE not in sys.path:
    sys.path.insert(0, BASE)
from monthly_reporting import classify_copq_type

DB_OVERVIEW = os.path.join(BASE, "Database", "CoPQ database 25.xlsx")
DB_TYPE     = os.path.join(BASE, "Database", "CoPQ_type_analysis.xlsx")
FTT_XLSX    = os.path.join(HERE, "FTT_Combined_Report.xlsx")
OUT         = os.path.join(HERE, "COPQ_Report_Data.xlsx")

MONTH_SHEET = "Aug-2026"          # overview sheet for the current period
MONTH_LABEL  = "Aug"              # column label in type_analysis rolling matrix
ROWS_SHEET  = "COPQ_rows_Aug-2026"

SITES = ["VH", "VH2", "VH3", "VH4", "JV", "JV2", "JVB"]
VN_SITES = ["VH", "VH2", "VH3", "VH4"]
ID_SITES = ["JV", "JV2", "JVB"]
TOUCHUP = "touch-up paint"
EXCLUDE = ["Rework"]

def bar_cat(row):
    group = classify_copq_type(row)
    mapping = {
        "B/C": "BC",
        "Touch-up": "TU",
        "Reinspection": "RE",
        "Rework": "RW",
    }
    return mapping.get(group)

def load_rows():
    try:
        df = pd.read_excel(DB_OVERVIEW, sheet_name=ROWS_SHEET, header=0)
    except ValueError:
        df = pd.DataFrame()
    return df

def main():
    wb = openpyxl.Workbook()
    # drop default sheet
    wb.remove(wb.active)

    rows = load_rows()
    if not rows.empty:
        rows = rows.copy()
        rows["cat"] = rows.apply(bar_cat, axis=1)

    # ---------- Slide1_Overview ----------
    ws = wb.create_sheet("Slide1_Overview")
    ws.append(["Site", "Touch up paint", "Re-inspection", "B/C Grade", "Total (excl Rework)"])
    tot = {"CLG": 0.0, "VN": 0.0, "ID": 0.0}
    for s in SITES:
        sub = rows[rows["Site"] == s] if not rows.empty else pd.DataFrame()
        def cost(cat):
            if sub.empty:
                return 0.0
            m = sub["cat"] == cat
            return round(float(sub[m]["Ttl cost ($)"].sum()), 2)
        tu, re, bc = cost("TU"), cost("RE"), cost("BC")
        ttl = round(tu + re + bc, 2)
        ws.append([s, tu, re, bc, ttl])
        tot["CLG"] += ttl
        if s in VN_SITES: tot["VN"] += ttl
        else: tot["ID"] += ttl
    ws.append(["CLG", None, None, None, round(tot["CLG"], 2)])
    ws.append([])
    ws.append(["Region", "Total (excl Rework)"])
    ws.append(["VN", round(tot["VN"], 2)])
    ws.append(["Indo", round(tot["ID"], 2)])
    ws.append(["CLG", round(tot["CLG"], 2)])

    # ---------- Slide2/3 by model (BC / RE / TU) ----------
    def by_model_sheet(title, cat):
        s = wb.create_sheet(title)
        s.append(["Site", "Model", "Cost ($)", "Qty (prs)"])
        if rows.empty:
            return
        for site in SITES:
            sub = rows[(rows["Site"] == site) & (rows["cat"] == cat)]
            if sub.empty:
                continue
            g = sub.groupby("Model").agg(
                cost=("Ttl cost ($)", "sum"), qty=("Defective Qty(Pair)", "sum")
            ).sort_values("cost", ascending=False).head(15)
            for model, r in g.iterrows():
                s.append([site, str(model).strip(), round(float(r["cost"]), 2),
                          round(float(r["qty"]), 1)])
        return s

    by_model_sheet("Slide2_BCGrade", "BC")
    by_model_sheet("Slide3_Reinspect", "RE")
    by_model_sheet("Slide3_Touchup", "TU")

    # ---------- Slide4_Monthly (rolling trend from type_analysis) ----------
    wm = wb.create_sheet("Slide4_Monthly")
    # read the Aug column + prior months to build a trend table
    ta = openpyxl.load_workbook(DB_TYPE, data_only=True)["Current month"]
    # header row 2 has month labels from col 3; find Aug
    cols = {}
    c = 3
    while c <= ta.max_column:
        v = ta.cell(2, c).value
        if v in (None, ""):
            break
        cols[str(v).strip()] = c
        c += 1
    aug_col = cols.get(MONTH_LABEL)
    if aug_col is None:
        if cols:
            fallback_label, aug_col = list(cols.items())[-1]
            print(f"[!] Warning: Month '{MONTH_LABEL}' not found in {DB_TYPE}['Current month']. "
                  f"Falling back to most recent month '{fallback_label}' (column {aug_col}).")
        else:
            raise ValueError(f"No valid month columns found in {DB_TYPE}['Current month'] header row 2.")
    # BC Grade cost rows: site rows under header row 2 (rows 3-6)
    bc_sites = ["VH", "VH2", "JV", "JV2"]
    re_sites = ["VH", "VH2", "VH3", "JV", "JV2", "JVB"]
    tu_sites = ["VH", "VH2", "VH4", "JV", "JVB"]
    wm.append(["Site", "Category", "Month", "Cost ($)"])
    def dump_block(sites, hdr_row, cat):
        for i, site in enumerate(sites):
            r = hdr_row + 1 + i
            for label, col in cols.items():
                v = ta.cell(r, col).value
                if isinstance(v, (int, float)):
                    wm.append([site, cat, label, round(float(v), 2)])
    dump_block(bc_sites, 2, "B/C Grade")
    dump_block(re_sites, 17, "Re-inspection")
    dump_block(tu_sites, 35, "Touch-up")

    # ---------- Slide4_TopDefect (FTT) ----------
    wd = wb.create_sheet("Slide4_TopDefect")
    ftt = pd.read_excel(FTT_XLSX, sheet_name="Top Defects Summary", header=1)
    # the summary sheet has site blocks; keep the meaningful columns
    keep = [c for c in ftt.columns if c and "Unnamed" not in str(c)]
    wd.append(list(keep))
    for _, r in ftt.iterrows():
        if pd.isna(r.iloc[0]) and all(pd.isna(r[c]) for c in keep):
            continue
        wd.append([r[c] for c in keep])

    # ---------- Slide4_Top3Model ----------
    wt = wb.create_sheet("Slide4_Top3Model")
    wt.append(["Site", "Model", "Cost ($)"])
    if not rows.empty:
        for site in SITES:
            sub = rows[(rows["Site"] == site) & (rows["cat"] != "RW")]
            if sub.empty:
                continue
            g = sub.groupby("Model")["Ttl cost ($)"].sum().sort_values(ascending=False).head(3)
            for model, v in g.items():
                wt.append([site, str(model).strip(), round(float(v), 2)])

    wb.save(OUT)
    print(f"[build_report_workbook] wrote {OUT}")
    print(f"  sheets: {wb.sheetnames}")

if __name__ == "__main__":
    main()
