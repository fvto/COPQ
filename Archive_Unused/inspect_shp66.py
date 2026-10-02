import pptx
from lxml import etree

prs = pptx.Presentation("Template_COPQ/Template_COPQ.pptx")
slide1 = prs.slides[1]

shp66 = [s for s in slide1.shapes if s.shape_id == 66][0]
print(f"Shape 66: name={shp66.name}, type={shp66.shape_type}")
print(etree.tostring(shp66._element, pretty_print=True).decode('utf-8')[:1000])
