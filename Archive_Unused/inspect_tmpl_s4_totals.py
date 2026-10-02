#!/usr/bin/env python3
"""Inspect the totals text and formatting on the template slide 4."""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = Presentation("Template_COPQ/Template_COPQ.pptx")
slide3 = prs.slides[3]  # Slide 4 (Touch-up)

for shape in slide3.shapes:
    if not shape.has_text_frame:
        continue
    text = shape.text_frame.text.strip()
    if "$" not in text:
        continue
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        continue
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.runs[0] if paragraph.runs else None
    font_name = run.font.name if run else None
    font_size = run.font.size.pt if run and run.font.size else None
    is_bold = run.font.bold if run else None
    font_color = run.font.color.rgb if run and run.font.color.type == 1 else None
    print(
        f"Slide 4 Shape '{shape.name}': text={text!r}, "
        f"font={font_name}, size={font_size}, bold={is_bold}, color={font_color}"
    )
