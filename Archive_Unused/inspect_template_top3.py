import pptx
from pptx.oxml.ns import qn

prs = pptx.Presentation("Template_COPQ/Template_COPQ.pptx")
slide1 = prs.slides[1]

defect_charts = ["Chart 78", "Chart 80", "Chart 92", "Chart 96"]
for name in defect_charts:
    shp = [s for s in slide1.shapes if s.name == name][0]
    ch = shp.chart
    print(f"\n=================== {name} (Template) ===================")
    print(f"Chart type: {ch.chart_type}")
    sers = list(ch._chartSpace.chart.plotArea.iter(qn('c:ser')))
    print(f"Number of series: {len(sers)}")
    
    # categories
    cat_names = []
    if sers:
        cat = sers[0].find(qn('c:cat'))
        if cat is not None:
            for pt in cat.iter(qn('c:pt')):
                v = pt.find(qn('c:v'))
                if v is not None: cat_names.append(v.text)
    print(f"Categories (X-axis): {cat_names}")
    
    # series names and colors
    for idx, s in enumerate(sers):
        tx = s.find(qn('c:tx'))
        v = tx.find('.//c:v', namespaces=s.nsmap) if tx is not None else None
        sname = v.text if v is not None else f"Ser {idx}"
        spPr = s.find(qn('c:spPr'))
        srgb = spPr.find('.//a:srgbClr', namespaces=s.nsmap) if spPr is not None else None
        col = srgb.get('val') if srgb is not None else "no-rgb"
        # values
        val_pts = []
        val = s.find(qn('c:val'))
        if val is not None:
            for pt in val.iter(qn('c:pt')):
                pv = pt.find(qn('c:v'))
                if pv is not None: val_pts.append(pv.text)
        print(f"  Ser {idx:2d}: {sname:30s} | color={col} | vals={val_pts}")
