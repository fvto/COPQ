import pptx
from pptx.oxml.ns import qn
from lxml import etree

prs = pptx.Presentation("Template_COPQ/Template_COPQ.pptx")
slide0 = prs.slides[0]

chart_names = ["Chart 18", "Chart 20", "Chart 22", "Chart 24", "Chart 26", "Chart 28", "Chart 30"]
for name in chart_names:
    shp = [s for s in slide0.shapes if s.name == name][0]
    ch = shp.chart
    for ser in ch._chartSpace.chart.plotArea.iter(qn('c:ser')):
        val = ser.find(qn('c:val'))
        if val is not None:
            nref = val.find(qn('c:numRef'))
            if nref is not None:
                vcache = nref.find(qn('c:numCache'))
                if vcache is not None:
                    fc = vcache.find(qn('c:formatCode'))
                    fc_text = fc.text if fc is not None else "None"
                    pts = [(pt.get('idx'), pt.get(qn('c:formatCode')), pt.find(qn('c:v')).text if pt.find(qn('c:v')) is not None else None) for pt in vcache.findall(qn('c:pt'))]
                    print(f"{name}: cache formatCode='{fc_text}', pts={pts}")
