import openpyxl
import pandas as pd

wb = openpyxl.load_workbook('Output/FTT_Combined_Report.xlsx', data_only=False)
ws = wb['Top Defects Summary']

sites = ['VH', 'VH2', 'JV', 'JV2']
site_data = {}

# Site blocks start where col A has the site name
current_site = None
header_row = None
for r in range(1, ws.max_row+1):
    val = ws.cell(r, 1).value
    if val in sites:
        current_site = val
        header_row = r
        defects = []
        c = 2
        while True:
            d_val = ws.cell(r, c).value
            if not d_val:
                break
            defects.append(d_val)
            c += 1
        site_data[current_site] = {
            'defects': defects,
            'models': []
        }
    elif current_site and r <= header_row + 3 and val:
        model_name = val
        rates = {}
        for c_idx, defect in enumerate(site_data[current_site]['defects'], start=2):
            cell = ws.cell(r, c_idx)
            formula = cell.value
            rates[defect] = formula
        site_data[current_site]['models'].append({
            'model': model_name,
            'rates': rates
        })

print("SUMMARY OF Top Defects Summary in FTT_Combined_Report.xlsx:")
for s, d in site_data.items():
    print(f"\n=== SITE {s} ===")
    print("Defects:", d['defects'])
    for m in d['models']:
        print(f"  Model: {m['model']}")
        for def_name, form in m['rates'].items():
            if form:
                print(f"    {def_name}: {form}")
