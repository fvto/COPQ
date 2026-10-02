import pptx

prs = pptx.Presentation("Template_COPQ/Template_COPQ.pptx")
slide1 = prs.slides[1]

for shp in slide1.shapes:
    print(f"ID {shp.shape_id:3d} | name={shp.name:25s} | type={shp.shape_type} | pos=({shp.left:8d}, {shp.top:8d}, {shp.width:8d}, {shp.height:8d})")
