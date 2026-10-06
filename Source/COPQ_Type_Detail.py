"""
COPQ_Type_Detail.py - Script COPQ Type Detail
=============================================
Python implementation of Output/COPQ_Type_Detail.ts.

Workflow Step 2:
- Targets sheet 'COPQ_Clean'.
- Adds column 'Type Detail'.
- For rows where Type is 'Reinspection':
    If Remark (Re-inspection Station) contains 'touch-up' or 'touch up',
    detail is classified as 'Touch-up'; otherwise 'Reinspection'.
- Updates table 'COPQ_CleanData' reference and columns if present.
"""

import os
import sys
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.table import TableColumn
from openpyxl.utils import get_column_letter

# Ensure UTF-8 stdout encoding for Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")

CLEAN_SHEET_NAME = "COPQ_Clean"
TABLE_NAME = "COPQ_CleanData"

F_TYPE = "Type"
F_REMARK_STATION = "Remark (Re-inspection Station)"
NEW_FIELD = "Type Detail"


def run_copq_type_detail(workbook_or_path):
    """
    Execute Script COPQ Type Detail.
    Accepts either an openpyxl.Workbook instance or a file path.
    Returns the modified openpyxl.Workbook.
    """
    is_path = isinstance(workbook_or_path, str)
    if is_path:
        wb = openpyxl.load_workbook(workbook_or_path)
    else:
        wb = workbook_or_path

    if CLEAN_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet '{CLEAN_SHEET_NAME}' not found.")

    ws = wb[CLEAN_SHEET_NAME]

    # Locate headers in row 1
    headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]

    try:
        type_idx = headers.index(F_TYPE) + 1
    except ValueError:
        raise ValueError(f"Column '{F_TYPE}' not found in '{CLEAN_SHEET_NAME}'.")

    try:
        remark_idx = headers.index(F_REMARK_STATION) + 1
    except ValueError:
        # Fallback to partial search if needed
        matching = [i + 1 for i, h in enumerate(headers) if "re-inspection station" in h.lower() or "remark" in h.lower()]
        if matching:
            remark_idx = matching[0]
        else:
            raise ValueError(f"Column '{F_REMARK_STATION}' not found in '{CLEAN_SHEET_NAME}'.")

    # Locate or create NEW_FIELD column
    if NEW_FIELD in headers:
        new_col_idx = headers.index(NEW_FIELD) + 1
    else:
        new_col_idx = len(headers) + 1
        headers.append(NEW_FIELD)
        hdr_cell = ws.cell(1, new_col_idx, NEW_FIELD)
        hdr_cell.fill = PatternFill(start_color="1F4E78", fill_type="solid")
        hdr_cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        hdr_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hdr_cell.border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

    data_font = Font(name="Calibri", size=9)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    count_touchup = 0
    count_reinsp = 0

    for r in range(2, ws.max_row + 1):
        type_val = str(ws.cell(r, type_idx).value or "").strip()
        remark_val = str(ws.cell(r, remark_idx).value or "").strip().lower()

        detail = type_val
        if type_val.lower() == "reinspection":
            if "touch-up" in remark_val or "touch up" in remark_val:
                detail = "Touch-up"
                count_touchup += 1
            else:
                detail = "Reinspection"
                count_reinsp += 1

        cell = ws.cell(r, new_col_idx, detail)
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="left", vertical="center")

    # Update Table reference if present
    for t_name in list(ws.tables.keys()):
        tab = ws.tables[t_name]
        max_row = ws.max_row
        new_col_letter = get_column_letter(max(new_col_idx, ws.max_column))
        tab.ref = f"A1:{new_col_letter}{max_row}"
        existing_col_names = [col.name for col in tab.tableColumns]
        if NEW_FIELD not in existing_col_names:
            tab.tableColumns.append(TableColumn(id=len(tab.tableColumns) + 1, name=NEW_FIELD))

    # Set column width
    ws.column_dimensions[get_column_letter(new_col_idx)].width = 18

    print(f"[+] Script COPQ Type Detail completed: Added/updated '{NEW_FIELD}' ({count_touchup} Touch-up, {count_reinsp} Reinspection).")

    if is_path:
        wb.save(workbook_or_path)
    return wb


if __name__ == "__main__":
    _dir = os.path.dirname(__file__)
    _base = os.path.dirname(_dir) if os.path.basename(_dir) == "Source" else _dir
    default_target = os.path.join(_base, "Output", "COPQ_Clean.xlsx")
    if not os.path.exists(default_target):
        default_target = os.path.join(_dir, "COPQ_Clean.xlsx")
    target = sys.argv[1] if len(sys.argv) > 1 else default_target
    if not os.path.exists(target):
        print(f"Target file not found: {target}")
        sys.exit(1)
    run_copq_type_detail(target)
