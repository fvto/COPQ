"""
Combined_COPQ.py - Script COPQ Clean
====================================
Python implementation of Output/Combined_COPQ.ts.

Workflow Step 1:
- Gathers data across all individual site sheets (excluding COPQ_Clean).
- Builds unified master header list.
- Validates data rows against quality and SOP rules.
- Generates COPQ_Clean master sheet with Table formatting, status columns,
  and highlights validation errors in soft red (#FFC7CE / #9C0006).
"""

import os
import sys
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

# Ensure UTF-8 stdout encoding for Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")

HOURS_FIELD = "Working hours"
QTY_FIELD = "Defective Qty(Pair)"
COST_FIELD = "Ttl cost ($)"
TYPE_FIELD = "Type"
CATEGORY_FIELD = "Category"
STYLE_FIELD = "Style Nbr"
MODEL_FIELD = "Model"

HOURS_LIMIT = 15   # error if hours > 15
QTY_LIMIT = 2000   # error if qty < 2000 (combined with hours > 15)
COST_LIMIT = 1000  # error if cost > 1000
REINSPECTION_TYPE = "reinspection"

REQUIRED_FIELDS = ["Date", "Category", "Type", "Factory", "Model"]
STYLE_FACTORIES = ["factory 1", "factory 2", "factory 3", "factory 4", "factory 5", "sandals", "factory sandals"]
MOLD_FACTORIES = ["ip", "os", "stockfitting"]


def _is_mold_factory(f_name):
    f = str(f_name or "").strip().lower()
    return any(k in f for k in ["ip", "os", "stockfitting", "sole", "bottom", "mold"])


def _is_style_factory(f_name):
    f = str(f_name or "").strip().lower()
    return any(k in f for k in ["factory", "fty", "sandals", "line", "stitch", "assembly", "upper"])


def _is_blank(val):
    if val is None:
        return True
    if isinstance(val, float) and (val != val):  # math.isnan check
        return True
    s = str(val).strip()
    return s == "" or s.lower() == "nan"


def _to_float(val):
    if _is_blank(val):
        return None
    try:
        s = str(val).replace("$", "").replace(",", "").strip()
        if s.lower() == "nan":
            return None
        return float(s)
    except (ValueError, TypeError):
        return None


def run_copq_clean(workbook_or_path):
    """
    Execute Script COPQ Clean.
    Accepts either an openpyxl.Workbook instance or a file path.
    Returns the modified openpyxl.Workbook.
    """
    is_path = isinstance(workbook_or_path, str)
    if is_path:
        wb = openpyxl.load_workbook(workbook_or_path)
    else:
        wb = workbook_or_path

    output_name = "COPQ_Clean"
    skip_sheets = [
        output_name,
        "COPQ_Pivot_Overview",
        "COPQ_Pivot_TopFactory",
        "COPQ_Pivot_TopModel",
        "COPQ_Pivot_ReInsp&TouchUp",
        "COPQ_Report",
    ]

    # Delete existing sheet if present
    if output_name in wb.sheetnames:
        del wb[output_name]

    out_ws = wb.create_sheet(title=output_name)

    # 1. Pass 1: Gather all unique headers across all site sheets and build reference maps
    master_headers = []
    seen_headers = set()
    style_to_model = {}
    mold_to_model = {}
    model_to_style = {}

    for sheet_name in wb.sheetnames:
        if sheet_name in skip_sheets:
            continue
        ws = wb[sheet_name]
        if ws.max_row < 3:
            continue

        # Row 2 contains column headers (Row 1 is Title "Cost Of Poor Quality")
        row_vals = [str(ws.cell(2, c).value or "").strip() for c in range(1, ws.max_column + 1)]
        for h in row_vals:
            if h and h not in seen_headers:
                seen_headers.add(h)
                master_headers.append(h)

        # Build cross-row reference maps for smart self-healing
        m_col_p1 = row_vals.index(MODEL_FIELD) if MODEL_FIELD in row_vals else -1
        s_col_p1 = row_vals.index(STYLE_FIELD) if STYLE_FIELD in row_vals else -1
        mold_col_p1 = row_vals.index("Mold") if "Mold" in row_vals else -1

        for r in range(3, ws.max_row + 1):
            m_val = str(ws.cell(r, m_col_p1 + 1).value or "").strip() if m_col_p1 >= 0 else ""
            if not m_val or m_val.lower() == "nan":
                continue
            if s_col_p1 >= 0:
                s_val = str(ws.cell(r, s_col_p1 + 1).value or "").strip()
                if s_val and s_val.lower() != "nan":
                    style_to_model[s_val.upper()] = m_val
                    model_to_style.setdefault(m_val.upper(), s_val)
            if mold_col_p1 >= 0:
                mold_val = str(ws.cell(r, mold_col_p1 + 1).value or "").strip()
                if mold_val and mold_val.lower() != "nan":
                    mold_to_model[mold_val.upper()] = m_val

    if not master_headers:
        print("[!] No valid site data found across sheets.")
        if is_path:
            wb.save(workbook_or_path)
        return wb

    full_headers = ["Site"] + master_headers + ["Missing Fields", "Validation Errors", "Data Quality Status"]
    combined_rows = []
    highlight_cells = []  # list of (row_idx_1_based, col_idx_1_based)

    # 2. Pass 2: Map row data dynamically into the uniform master header structure
    for sheet_name in wb.sheetnames:
        if sheet_name in skip_sheets:
            continue
        ws = wb[sheet_name]
        if ws.max_row < 3:
            continue

        raw_sheet_headers = [str(ws.cell(2, c).value or "").strip() for c in range(1, ws.max_column + 1)]

        def get_col_idx(field_name):
            try:
                return raw_sheet_headers.index(field_name)
            except ValueError:
                return -1

        hours_idx = get_col_idx(HOURS_FIELD)
        qty_idx = get_col_idx(QTY_FIELD)
        cost_idx = get_col_idx(COST_FIELD)
        type_idx = get_col_idx(TYPE_FIELD)
        category_idx = get_col_idx(CATEGORY_FIELD)
        style_idx = get_col_idx(STYLE_FIELD)
        model_idx = get_col_idx(MODEL_FIELD)
        factory_idx = get_col_idx("Factory")
        mold_idx = get_col_idx("Mold")
        manpower_idx = get_col_idx("Manpower")
        sgs_qty_idx = get_col_idx("SGS Retest Qty (Pic)")
        sgs_price_idx = get_col_idx("SGS Test unit price/pic ($)")

        for r in range(3, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if all(v is None or str(v).strip() == "" for v in row_vals):
                continue

            row_map = {}
            for i, h in enumerate(raw_sheet_headers):
                if h and i < len(row_vals):
                    row_map[h] = row_vals[i]

            # Exclude Rework records per user requirement
            type_raw = str(row_map.get(TYPE_FIELD, "") or "").strip().lower()
            if type_raw == "rework":
                continue

            normalized = ["" if _is_blank(row_map.get(h)) else row_map.get(h) for h in master_headers]

            missing = [
                h for h in REQUIRED_FIELDS
                if h not in row_map or _is_blank(row_map[h])
            ]

            # --- Validation checks ---
            validation_errors = []
            local_highlights = []

            def get_val(idx):
                return row_vals[idx] if 0 <= idx < len(row_vals) else None

            hours_val = _to_float(get_val(hours_idx))
            qty_val = _to_float(get_val(qty_idx))
            cost_val = _to_float(get_val(cost_idx))
            manpower_val = _to_float(get_val(manpower_idx))
            sgs_qty_val = _to_float(get_val(sgs_qty_idx))
            sgs_price_val = _to_float(get_val(sgs_price_idx))

            has_hours = hours_val is not None
            has_qty = qty_val is not None
            has_cost = cost_val is not None

            # Negative numbers anomaly check
            if has_cost and cost_val < 0:
                validation_errors.append(f"Negative {COST_FIELD}")
                local_highlights.append(COST_FIELD)
            if has_qty and qty_val < 0:
                validation_errors.append(f"Negative {QTY_FIELD}")
                local_highlights.append(QTY_FIELD)
            if has_hours and hours_val < 0:
                validation_errors.append(f"Negative {HOURS_FIELD}")
                local_highlights.append(HOURS_FIELD)

            # Rule 1A: Extreme Working hours > 15 AND Defective Qty(Pair) < 2000
            if has_hours and has_qty and hours_val > HOURS_LIMIT and qty_val < QTY_LIMIT:
                validation_errors.append(f"{HOURS_FIELD} > {HOURS_LIMIT} & {QTY_FIELD} < {QTY_LIMIT}")
                local_highlights.extend([HOURS_FIELD, QTY_FIELD])

            # Rule 1B (Slide 1 VH3): Working hours > 5 AND Defective Qty(Pair) < 300
            elif has_hours and has_qty and hours_val > 5.0 and qty_val < 300:
                validation_errors.append(f"{HOURS_FIELD} > 5 & {QTY_FIELD} < 300")
                local_highlights.extend([HOURS_FIELD, QTY_FIELD])

            # Rule 1C: Abnormal inspection rate (< 25 pairs/h for extended hours > 3h)
            elif has_hours and has_qty and hours_val > 3.0 and qty_val > 0 and (qty_val / hours_val) < 25.0:
                validation_errors.append("Abnormal inspection rate (< 25 pairs/h)")
                local_highlights.extend([HOURS_FIELD, QTY_FIELD])

            # Rule 2: Ttl cost ($) > 1000
            if has_cost and cost_val > COST_LIMIT:
                validation_errors.append(f"{COST_FIELD} > {COST_LIMIT}")
                local_highlights.append(COST_FIELD)

            type_raw = str(get_val(type_idx) or "").strip().lower()
            is_reinspection = (type_raw == REINSPECTION_TYPE)
            is_bc_grade = ("b/c" in type_raw) or ("b grade" in type_raw) or ("c grade" in type_raw)

            factory_raw = str(get_val(factory_idx) or "").strip().lower()
            mold_raw = get_val(mold_idx)
            mold_blank = (mold_idx >= 0 and _is_blank(mold_raw))

            uses_mold = _is_mold_factory(factory_raw)
            uses_style_nbr = not uses_mold and _is_style_factory(factory_raw)

            category_raw = get_val(category_idx)
            category_blank = (category_idx >= 0 and _is_blank(category_raw))

            type_blank = (type_idx >= 0 and _is_blank(type_raw))

            style_raw = get_val(style_idx)
            style_blank = (style_idx >= 0 and _is_blank(style_raw))

            model_raw = get_val(model_idx)
            model_blank = (model_idx >= 0 and _is_blank(model_raw))

            # Smart Self-Healing for missing Model / Style Nbr via cross-row lookup
            auto_resolved_notes = []
            if model_blank:
                if not style_blank and str(style_raw).strip().upper() in style_to_model:
                    resolved_m = style_to_model[str(style_raw).strip().upper()]
                    model_raw = resolved_m
                    model_blank = False
                    if MODEL_FIELD in master_headers:
                        normalized[master_headers.index(MODEL_FIELD)] = resolved_m
                    row_map[MODEL_FIELD] = resolved_m
                    auto_resolved_notes.append(f"Auto-resolved Model from Style Nbr: {resolved_m}")
                elif not mold_blank and str(mold_raw).strip().upper() in mold_to_model:
                    resolved_m = mold_to_model[str(mold_raw).strip().upper()]
                    model_raw = resolved_m
                    model_blank = False
                    if MODEL_FIELD in master_headers:
                        normalized[master_headers.index(MODEL_FIELD)] = resolved_m
                    row_map[MODEL_FIELD] = resolved_m
                    auto_resolved_notes.append(f"Auto-resolved Model from Mold: {resolved_m}")

            if style_blank and uses_style_nbr and not model_blank:
                if str(model_raw).strip().upper() in model_to_style:
                    resolved_s = model_to_style[str(model_raw).strip().upper()]
                    style_raw = resolved_s
                    style_blank = False
                    if STYLE_FIELD in master_headers:
                        normalized[master_headers.index(STYLE_FIELD)] = resolved_s
                    row_map[STYLE_FIELD] = resolved_s
                    auto_resolved_notes.append(f"Auto-resolved Style Nbr from Model: {resolved_s}")

            # Re-evaluate missing fields after potential self-healing
            missing = [
                h for h in REQUIRED_FIELDS
                if h not in row_map or _is_blank(row_map[h])
            ]

            hours_blank_or_zero = (hours_idx >= 0 and (hours_val is None or hours_val == 0))
            qty_blank_or_zero = (qty_idx >= 0 and (qty_val is None or qty_val == 0))
            cost_blank_or_zero = (cost_idx >= 0 and (cost_val is None or cost_val == 0))

            manpower_raw = get_val(manpower_idx)
            manpower_blank_or_zero = (
                manpower_idx >= 0 and (_is_blank(manpower_raw) or manpower_val is None or manpower_val == 0)
            )

            # Rule 3: Missing data (Slide 3 JV2)
            # Row where key metrics (model, qty, hours, cost, manpower) are empty or zero
            is_missing_data = (
                model_blank and qty_blank_or_zero and hours_blank_or_zero and cost_blank_or_zero and manpower_blank_or_zero
            )

            if is_missing_data:
                validation_errors.append("Missing data")
                local_highlights.extend([
                    MODEL_FIELD,
                    QTY_FIELD,
                    "Manpower",
                    HOURS_FIELD,
                    COST_FIELD
                ])
                if uses_mold:
                    local_highlights.append("Mold")
                elif uses_style_nbr:
                    local_highlights.append(STYLE_FIELD)
            else:
                # Rule 4: Slide 3 JV2 Row 3 - Wrong Working hours => Wrong Ttl cost ($)
                if is_reinspection and hours_blank_or_zero and not qty_blank_or_zero:
                    validation_errors.append(f"Wrong {HOURS_FIELD} => Wrong {COST_FIELD}")
                    local_highlights.extend([HOURS_FIELD, COST_FIELD])
                    if manpower_val and manpower_val > 0:
                        local_highlights.append("Manpower")

                if is_reinspection and qty_blank_or_zero and not hours_blank_or_zero:
                    validation_errors.append(f"Wrong {QTY_FIELD}")
                    local_highlights.append(QTY_FIELD)

                if is_reinspection and not hours_blank_or_zero and cost_blank_or_zero:
                    validation_errors.append(f"Working hours present but zero {COST_FIELD}")
                    local_highlights.extend([HOURS_FIELD, COST_FIELD])

                # Rule 5: Slide 2 VH - Model Linking Rules
                if is_reinspection:
                    if uses_mold:
                        if mold_blank:
                            # Slide 2 Row 1: Thiếu mold code => không link model
                            validation_errors.append("Missing mold code => cannot link model")
                            local_highlights.extend(["Mold", MODEL_FIELD])
                        elif model_blank:
                            # Slide 2 Row 2 & 3: Sai Mold code => Không link được model
                            validation_errors.append("Invalid mold code => cannot link model")
                            local_highlights.extend(["Mold", MODEL_FIELD])
                    elif uses_style_nbr:
                        if style_blank:
                            # Slide 2 Row 4: Thiếu hình thể => không link được model
                            validation_errors.append(f"Missing {STYLE_FIELD.lower()} => cannot link model")
                            local_highlights.extend([STYLE_FIELD, MODEL_FIELD])
                        elif model_blank:
                            # Sai hình thể => Không link được model
                            validation_errors.append(f"Invalid {STYLE_FIELD.lower()} => cannot link model")
                            local_highlights.extend([STYLE_FIELD, MODEL_FIELD])
                    else:
                        if model_blank:
                            validation_errors.append("Missing model => cannot link model")
                            local_highlights.append(MODEL_FIELD)

                # B/C Grade consistency checks
                if is_bc_grade:
                    if qty_blank_or_zero and not cost_blank_or_zero:
                        validation_errors.append(f"B/C Grade missing {QTY_FIELD}")
                        local_highlights.append(QTY_FIELD)
                    elif cost_blank_or_zero and not qty_blank_or_zero:
                        validation_errors.append(f"B/C Grade missing {COST_FIELD}")
                        local_highlights.append(COST_FIELD)

                # SGS Retest consistency checks
                if sgs_qty_val and sgs_qty_val > 0 and (sgs_price_val is None or sgs_price_val == 0):
                    validation_errors.append("SGS Retest Qty recorded without unit price")
                    local_highlights.extend(["SGS Retest Qty (Pic)", "SGS Test unit price/pic ($)"])
                elif sgs_price_val and sgs_price_val > 0 and (sgs_qty_val is None or sgs_qty_val == 0):
                    validation_errors.append("SGS unit price recorded without Retest Qty")
                    local_highlights.extend(["SGS Retest Qty (Pic)", "SGS Test unit price/pic ($)"])

                if type_blank and cost_blank_or_zero:
                    validation_errors.append(f"Blank {TYPE_FIELD} => Blank {COST_FIELD}")
                    local_highlights.extend([TYPE_FIELD, COST_FIELD])
                elif type_blank:
                    validation_errors.append(f"{TYPE_FIELD} blank")
                    local_highlights.append(TYPE_FIELD)

                if category_blank:
                    validation_errors.append(f"{CATEGORY_FIELD} blank")
                    local_highlights.append(CATEGORY_FIELD)

            # Add auto-resolved notices to validation_errors
            validation_errors.extend(auto_resolved_notes)

            # Data Quality Status
            actual_errors = [e for e in validation_errors if not e.startswith("Auto-resolved")]
            if missing and actual_errors:
                status = "Missing Data & Validation Error"
            elif missing:
                status = "Missing Data"
            elif actual_errors:
                status = "Validation Error"
            elif auto_resolved_notes:
                status = "Auto-Resolved"
            else:
                status = "OK"

            combined_rows.append([
                sheet_name,
                *normalized,
                ", ".join(missing),
                ", ".join(validation_errors),
                status
            ])

            row_idx_in_combined = len(combined_rows) + 1  # 1-based, header is row 1
            for field in set(local_highlights):
                if field in master_headers:
                    col_idx_in_combined = 2 + master_headers.index(field)  # Col 1 is Site
                    highlight_cells.append((row_idx_in_combined, col_idx_in_combined))

    if not combined_rows:
        print("[!] No data rows to combine.")
        if is_path:
            wb.save(workbook_or_path)
        return wb

    # Write headers and rows to COPQ_Clean
    out_ws.append(full_headers)
    for row in combined_rows:
        out_ws.append(row)

    # Style Header
    hdr_fill = PatternFill(start_color="1F4E78", fill_type="solid")
    hdr_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for c in range(1, len(full_headers) + 1):
        cell = out_ws.cell(1, c)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Add Table
    max_row = len(combined_rows) + 1
    max_col_letter = get_column_letter(len(full_headers))
    table = Table(displayName="COPQ_CleanData", ref=f"A1:{max_col_letter}{max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )
    out_ws.add_table(table)

    # Apply highlight styling to failing cells
    fail_fill = PatternFill(start_color="FFC7CE", fill_type="solid")
    fail_font = Font(name="Calibri", size=9, color="9C0006")

    for r_idx, c_idx in highlight_cells:
        cell = out_ws.cell(r_idx, c_idx)
        cell.fill = fail_fill
        cell.font = fail_font

    # Freeze panes below header
    out_ws.freeze_panes = "A2"

    # Auto-fit column widths
    for col in out_ws.columns:
        max_len = 10
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)), 45))
        col_letter = get_column_letter(col[0].column)
        out_ws.column_dimensions[col_letter].width = max_len + 3

    print(f"[+] Script COPQ Clean completed: {len(combined_rows)} rows combined on 'COPQ_Clean' with {len(highlight_cells)} highlighted issues.")

    if is_path:
        wb.save(workbook_or_path)
    return wb


if __name__ == "__main__":
    _dir = os.path.dirname(__file__)
    _base = os.path.dirname(_dir) if os.path.basename(_dir) == "Source" else _dir
    default_target = os.path.join(_base, "Output", "COPQ_Clean.xlsx")
    if not os.path.exists(default_target):
        default_target = os.path.join(_dir, "COPQ_Clean.xlsx")
    target = sys.argv[1] if len(sys.argv) > 1 else default_target
    if not os.path.exists(target):
        print(f"Target file not found: {target}")
        sys.exit(1)
    run_copq_clean(target)
