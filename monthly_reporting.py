"""Month-aware CoPQ reporting utilities shared by the dashboard and exporters.

The helpers here implement the monthly archive rules in the CoPQ SOP while
leaving the two source workbooks in ``Database`` untouched.  Each run creates
working clones under Output/Monthly/<YYYY-MM>/Database.
"""

from __future__ import annotations

import datetime as dt
import os
import re
import shutil
from pathlib import Path

import openpyxl
import pandas as pd


SITES = ("VH", "VH2", "VH3", "VH4", "JV", "JV2", "JVB")


def report_month_from_inputs(input_dir: str, fallback: dt.date | None = None) -> dt.date:
    """Read YYYY-MM-DD date ranges in downloaded ERP filenames.

    A mixed-month directory is rejected so a monthly report cannot silently
    combine periods.  If no recognisable filename is present, use ``fallback``
    (or the current month) so manually named test exports remain usable.
    """
    months = set()
    for name in os.listdir(input_dir) if os.path.isdir(input_dir) else ():
        match = re.search(r"(20\d{2})-(\d{2})-\d{2}", name)
        if match:
            try:
                year, month = int(match.group(1)), int(match.group(2))
                if 1 <= month <= 12:
                    months.add((year, month))
                else:
                    print(f"[!] Warning: Ignoring filename with invalid month ({month}) in '{name}'")
            except (ValueError, TypeError):
                print(f"[!] Warning: Could not parse date from filename '{name}'")
    if len(months) > 1:
        raise ValueError("Input folder contains exports from more than one report month.")
    if months:
        year, month = months.pop()
        return dt.date(year, month, 1)
    fallback = fallback or dt.date.today()
    return dt.date(fallback.year, fallback.month, 1)


def month_sheet_name(month: dt.date) -> str:
    return month.strftime("%B-%Y")


def type_analysis_sheet_name(month: dt.date) -> str:
    return month.strftime("%b-%y")


def month_label(month: dt.date) -> str:
    return month.strftime("%b, %Y")


def classify_copq_type(row: pd.Series) -> str:
    """Apply the SOP: Touch-up Paint is a subset of Reinspection by remark."""
    raw_type = str(row.get("Type", "")).strip().lower()
    remark = str(row.get("Remark (Re-inspection Station)", "")).strip().lower()
    if "touch" in raw_type or "touch-up" in remark or "touch up" in remark:
        return "Touch-up"
    if raw_type in {"b/c", "bc"}:
        return "B/C"
    if "reinspect" in raw_type or "re-inspection" in raw_type:
        return "Reinspection"
    if "rework" in raw_type:
        return "Rework"
    return "Other"


def validate_sop_data(master: pd.DataFrame) -> list[str]:
    """Return concise, actionable data-quality exceptions from SOP pages 7-8."""
    issues: list[str] = []
    if master.empty:
        return ["No CoPQ records were found after import."]
    working = master.copy()
    working["_group"] = working.apply(classify_copq_type, axis=1)
    if (working["_group"] == "Other").any():
        issues.append(f"{int((working['_group'] == 'Other').sum())} row(s) have an unsupported or blank Type.")
    for group in ("B/C", "Reinspection"):
        subset = working[working["_group"] == group]
        if subset.empty:
            continue
        required = ["Defective Qty(Pair)", "Ttl cost ($)"]
        if group == "Reinspection":
            required.extend(["Manpower", "Working hours"])
        for column in required:
            if column in subset.columns:
                blanks = pd.to_numeric(subset[column], errors="coerce").isna().sum()
                if blanks:
                    issues.append(f"{group}: {int(blanks)} row(s) missing {column}.")
    return issues


def _write_overview(ws, report: pd.DataFrame, month: dt.date) -> None:
    ws["A1"] = f"COPQ Overview (Based on {month.strftime('%b %Y')})"
    rows = {str(row["Factory"]).strip(): row for _, row in report.iterrows()}
    for excel_row, site in enumerate(SITES, start=3):
        row = rows.get(site, {})
        ws.cell(excel_row, 1).value = site
        ws.cell(excel_row, 2).value = float(row.get("Touch up paint", 0) or 0) or None
        ws.cell(excel_row, 3).value = float(row.get("Re-inspection", 0) or 0) or None
        ws.cell(excel_row, 4).value = float(row.get("B/C Grade", 0) or 0) or None
        ws.cell(excel_row, 5).value = float(row.get("Rework", 0) or 0) or None
        ws.cell(excel_row, 6).value = f"=SUM(B{excel_row}:E{excel_row})"
        ws.cell(excel_row, 7).value = f"=B{excel_row}+C{excel_row}"
    ws["A10"] = "CLG"
    for col in range(2, 8):
        letter = openpyxl.utils.get_column_letter(col)
        ws.cell(10, col).value = f"=SUM({letter}3:{letter}9)"
    ws["J3"] = "=SUM(F3:F6)"
    ws["M3"] = "=SUM(F7:F9)"
    ws["K6"] = "=J3"
    ws["K7"] = "=M3"
    ws["K8"] = "=SUM(K6:K7)"


def _report_from_clean_master(report_path: str) -> pd.DataFrame:
    """Build the internal monthly overview without requiring a report sheet."""
    master = pd.read_excel(report_path, sheet_name="COPQ_Clean")
    if master.empty or "Site" not in master.columns or "Ttl cost ($)" not in master.columns:
        return pd.DataFrame(columns=["Factory", "Touch up paint", "Re-inspection", "B/C Grade", "Rework"])
    master = master.copy()
    master["Factory"] = master["Site"].astype(str).str.strip().str.upper().replace({"JV3": "JVB"})
    master["_cost"] = pd.to_numeric(master["Ttl cost ($)"], errors="coerce").fillna(0.0)
    master["_type"] = master.apply(classify_copq_type, axis=1)
    type_columns = {
        "Touch-up": "Touch up paint",
        "Reinspection": "Re-inspection",
        "B/C": "B/C Grade",
        "Rework": "Rework",
    }
    report = master[["Factory"]].drop_duplicates().set_index("Factory")
    for source_type, output_column in type_columns.items():
        values = master.loc[master["_type"] == source_type].groupby("Factory")["_cost"].sum()
        report[output_column] = values
    return report.fillna(0.0).reset_index()


def create_monthly_database_clones(database_dir: str, output_dir: str, report_path: str, month: dt.date) -> dict[str, str]:
    """Clone both database files and add a formatted sheet for the report month."""
    destination = Path(output_dir) / "Monthly" / month.strftime("%Y-%m") / "Database"
    destination.mkdir(parents=True, exist_ok=True)
    report = _report_from_clean_master(report_path)
    result: dict[str, str] = {}
    for filename, source_sheet, target_sheet in (
        ("CoPQ database 25.xlsx", "Current", month_sheet_name(month)),
        ("CoPQ_type_analysis.xlsx", "Current month", type_analysis_sheet_name(month)),
    ):
        source = Path(database_dir) / filename
        target = destination / filename
        shutil.copy2(source, target)
        wb = openpyxl.load_workbook(target)
        if target_sheet in wb.sheetnames:
            del wb[target_sheet]
        ws = wb.copy_worksheet(wb[source_sheet])
        ws.title = target_sheet
        if filename.startswith("CoPQ database"):
            _write_overview(ws, report, month)
        else:
            ws["A1"] = f"CoPQ Type Analysis - {month_label(month)}"
        wb.save(target)
        result[filename] = str(target)
    return result
